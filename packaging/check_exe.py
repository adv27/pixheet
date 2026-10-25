"""Test a built Pixheet app: it must start, convert an image, and give each cell its color.

The app is dist/Pixheet.exe on Windows and dist/Pixheet on Linux. The script works on both
systems. On Linux, it needs an X server, for example Xvfb with xvfb-run.

The script does these steps:

1. It runs the app with ``--smoke``, one time with the light theme and one time with the
   dark theme. The exit code must be 0. The app must name the Sun Valley theme that it
   applied, not a fallback theme.
2. It makes a test image in a temporary folder and runs the app with ``--smoke`` and the
   image path. The app converts the image and closes. The script reads the .xlsx file
   with openpyxl and compares each cell with its pixel.
3. It makes one image for each of many formats and modes, and the app converts them all.
   The sheet and the styles of each .xlsx file must be the same, byte for byte, as the
   output of ``pixheet.convert`` in this Python. The spec leaves many modules out of the
   app. If the app needs one of them, this step fails.

If the app does not stop in time, the script stops the app and all processes that the
app started. At the end, the script deletes the folder.

Each run of the app has its own temporary folder: the script sets TEMP, TMP and TMPDIR.
A one-file app extracts itself to a folder ``_MEI...`` in that folder. After each run,
also after the script stops the app, the script deletes the folder. Thus a stopped app
cannot leave a ``_MEI...`` folder in the temporary folder of the user.

Each run of the app also gets ``PIXHEET_HIDDEN_WINDOW=1``. Then the window of the app
opens off the screen and does not take the focus. Thus the check shows no window on the
desktop of the person who runs the build.

Run the script from the repository root:

    uv run python packaging/check_exe.py dist/Pixheet.exe
    xvfb-run -a uv run python packaging/check_exe.py dist/Pixheet
"""

import argparse
import io
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

import openpyxl
from PIL import ExifTags, Image, ImageCms

import pixheet

WIDTH, HEIGHT = 48, 32
TIMEOUT = 120
"""The maximum time for the app, in seconds. The first start of a one-file app is slow."""
KILL_WAIT = 10
"""The maximum time to wait for the output after the script stops the app, in seconds."""
REMOVE_WAIT = 10
"""The maximum time to try again to delete the temporary folder of the app, in seconds."""
TEMP_VARIABLES = ("TEMP", "TMP", "TMPDIR")
"""The environment variables that give the temporary folder. Windows reads TMP and TEMP."""
MEMBERS = ("xl/worksheets/sheet1.xml", "xl/styles.xml")
"""The parts of an .xlsx file that the format check compares."""
FORMATS_FOLDER = "formats ảnh"
"""The folder of the format images. The name has a space and Vietnamese letters."""
THEME_ENV = "PIXHEET_THEME"
"""The environment variable that sets the theme of the app."""
HIDDEN_ENV = "PIXHEET_HIDDEN_WINDOW"
"""The environment variable that keeps the window of the app off the screen."""
THEME_LINE = re.compile(r"^pixheet-gui: theme (\S+)\s*$", re.MULTILINE)
"""The line of ``--smoke`` that names the ttk theme of the window."""
DND_LINE = re.compile(r"^pixheet-gui: drag and drop (on|off)\s*$", re.MULTILINE)
"""The line of ``--smoke`` that tells if the window has drag and drop.

The app itself fails the smoke test when it has no drag and drop but tkdnd can load. On
Linux with a libX11 older than 1.7, tkdnd cannot load, and "off" is correct.
"""


def make_image(path: Path) -> Image.Image:
    """Save a test image with transparent pixels and a few colors. Return its RGB pixels."""
    image = Image.new("RGBA", (WIDTH, HEIGHT))
    for y in range(HEIGHT):
        for x in range(WIDTH):
            alpha = 0 if (x + y) % 7 == 0 else 255
            image.putpixel((x, y), (x * 5, y * 7, (x * y) % 256, alpha))
    image.save(path)
    background = Image.new("RGBA", image.size, (255, 255, 255, 255))
    return Image.alpha_composite(background, image).convert("RGB")


def gradient(width: int, height: int) -> Image.Image:
    """Return an RGB image with many colors."""
    data = bytes(
        channel
        for y in range(height)
        for x in range(width)
        for channel in (x * 255 // (width - 1), y * 255 // (height - 1), (x * y) % 256)
    )
    return Image.frombytes("RGB", (width, height), data)


def format_cases() -> dict[str, tuple[Image.Image, dict]]:
    """Return the images of the format check, and the options to save them.

    The key is the file name. Each case uses a different decoder, or a different path in
    Pillow or in pixheet.
    """
    rgb = gradient(40, 30)
    rgba = rgb.copy()
    rgba.putalpha(Image.linear_gradient("L").resize(rgb.size))
    words = ((x * 1600 + y * 70) & 0xFFFF for y in range(30) for x in range(40))
    i16 = Image.frombytes("I;16", rgb.size, b"".join(w.to_bytes(2, "little") for w in words))
    srgb = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
    exif = Image.Exif()
    exif[ExifTags.Base.Orientation] = 6
    return {
        "rgba.png": (rgba, {}),
        "palette transparent.png": (rgb.quantize(16), {"transparency": 0}),
        "i16.png": (i16, {}),
        "cmyk icc.jpg": (rgb.convert("CMYK"), {"icc_profile": srgb, "quality": 95}),
        "exif rotated.jpg": (rgb, {"exif": exif, "quality": 95}),
        "rgba.webp": (rgba, {"lossless": True}),
        "rgba.avif": (rgba, {"quality": 90}),
        "palette.gif": (rgb.quantize(64), {}),
        "lzw.tif": (rgb, {"compression": "tiff_lzw"}),
        "lab.tif": (rgb.convert("LAB"), {}),
        "rgb.bmp": (rgb, {}),
    }


def sheet_parts(data: bytes) -> dict[str, bytes]:
    """Return the parts in ``MEMBERS`` of the .xlsx file ``data``."""
    with zipfile.ZipFile(io.BytesIO(data)) as xlsx:
        return {member: xlsx.read(member) for member in MEMBERS}


def make_format_images(folder: Path) -> tuple[dict[Path, dict[str, bytes]], list[str]]:
    """Save the format images in ``folder``, and convert each one with ``pixheet.convert``.

    Return the parts of each expected .xlsx file, and the problems. Each image has a
    different stem, because the app writes the .xlsx file next to the image.
    """
    folder.mkdir(parents=True, exist_ok=True)
    expected = {}
    problems = []
    for name, (image, options) in format_cases().items():
        stem, suffix = name.rsplit(".", 1)
        path = folder / f"{stem} {suffix}.{suffix}"
        try:
            image.save(path, **options)
            output = io.BytesIO()
            pixheet.convert(path, output)
        except Exception as e:  # each failure is a problem of the check
            problems.append(f"This Python cannot make the test file {name}: {e!r}")
            continue
        expected[path] = sheet_parts(output.getvalue())
    return expected, problems


def output_problems(expected: dict[Path, dict[str, bytes]]) -> list[str]:
    """Compare the .xlsx file of each image in ``expected`` with the expected parts."""
    problems = []
    for path, parts in expected.items():
        out = path.with_suffix(".xlsx")
        if not out.is_file():
            problems.append(f"The app did not write {out.name}.")
            continue
        got = sheet_parts(out.read_bytes())
        if different := [m for m in MEMBERS if got[m] != parts[m]]:
            problems.append(f"The app wrote a different {', '.join(different)} for {path.name}.")
    return problems


def check_formats(exe: Path, folder: Path, timeout: float = TIMEOUT) -> list[str]:
    """Convert the format images in ``folder`` with ``exe`` in one run. Return the problems."""
    expected, problems = make_format_images(folder)
    if problems:
        return problems
    start = time.monotonic()
    code, output = run_exe([str(exe), "--smoke", *map(str, expected)], timeout)
    seconds = time.monotonic() - start
    if problems := result_problems(code, output, timeout):
        return problems
    if problems := output_problems(expected):
        return problems
    print(f"The app converted {len(expected)} images of different formats in {seconds:.1f} s.")
    return []


def kill_tree(process: subprocess.Popen[bytes]) -> None:
    """Stop ``process`` and all processes that it started.

    A one-file app is two processes: the bootloader and the app. When the app shows an
    error dialog, it does not stop. If the script stops only the bootloader, the app
    keeps the app locked and keeps the output pipe open.
    """
    if sys.platform == "win32":
        subprocess.run(
            ["taskkill", "/PID", str(process.pid), "/T", "/F"], capture_output=True, check=False
        )
    else:
        os.killpg(process.pid, signal.SIGKILL)
    process.kill()


class CleanupError(Exception):
    """The script cannot delete the temporary folder of a run of the app."""


def remove_folder(folder: Path, wait: float = REMOVE_WAIT) -> None:
    """Delete ``folder`` and all its files. Raise ``CleanupError`` if that is not possible.

    After Windows stops a process, its files can stay locked for a short time. Thus the
    function tries again until ``wait`` seconds pass.
    """
    deadline = time.monotonic() + wait
    while True:
        try:
            shutil.rmtree(folder)
            return
        except FileNotFoundError:
            return
        except OSError as e:
            if time.monotonic() >= deadline:
                raise CleanupError(f"The script cannot delete the folder {folder}: {e}") from e
        time.sleep(0.2)


def run_exe(
    command: list[str], timeout: float = TIMEOUT, env: dict[str, str] | None = None
) -> tuple[int | None, str]:
    """Run ``command``, and return its exit code and its output (stdout and stderr).

    ``env`` adds environment variables. The command gets a new temporary folder in TEMP,
    TMP and TMPDIR, and ``PIXHEET_HIDDEN_WINDOW=1``. If the command does not stop in
    ``timeout`` seconds, stop it and all processes that it started. Then the exit code is
    None. At the end, delete the temporary folder. Raise ``CleanupError`` if that is not
    possible.
    """
    temp = Path(tempfile.mkdtemp(prefix="pixheet-exe-"))
    try:
        variables = dict.fromkeys(TEMP_VARIABLES, str(temp)) | {HIDDEN_ENV: "1"} | (env or {})
        return run_process(command, timeout, {**os.environ, **variables})
    finally:
        remove_folder(temp)


def run_process(command: list[str], timeout: float, env: dict[str, str]) -> tuple[int | None, str]:
    """Run ``command`` with the environment ``env``. ``run_exe`` tells the result."""
    # On POSIX, a new session lets kill_tree stop the child processes too.
    new_session = sys.platform != "win32"
    process = subprocess.Popen(
        command,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        start_new_session=new_session,
        env=env,
    )
    try:
        # communicate() closes the pipe and waits for the process.
        output, _ = process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        kill_tree(process)
        try:
            output, _ = process.communicate(timeout=KILL_WAIT)
        except subprocess.TimeoutExpired:
            # A process that kill_tree did not stop keeps the pipe open. Do not close the
            # pipe here, because the close waits until a read in a different thread ends.
            process.wait()
            output = b""
        return None, output.decode("utf-8", errors="replace")
    return process.returncode, output.decode("utf-8", errors="replace")


def result_problems(code: int | None, output: str, timeout: float) -> list[str]:
    """Return the problems of a run of the app with the exit code ``code``.

    A problem includes the output of the app, because the output tells the cause, for
    example "cannot open a window".
    """
    if code is None:
        problem = f"The app did not stop in {timeout:g} s. The script stopped it."
    elif code != 0:
        problem = f"The app stopped with exit code {code}."
    else:
        return []
    return with_output(problem, output)


def with_output(problem: str, output: str) -> list[str]:
    """Return ``problem`` with the output of the app, because the output tells the cause."""
    output = output.strip()
    return [f"{problem} The output of the app:\n{output}" if output else problem]


def theme_problems(output: str, mode: str) -> list[str]:
    """Return a problem if the output of ``--smoke`` does not name the theme ``mode``.

    ``mode`` is "light" or "dark". The app prints the ttk theme of its window. A
    different name, for example "vista", tells that the Sun Valley theme did not load.
    """
    match = THEME_LINE.search(output)
    if match is None:
        problem = "The app did not name its theme."
    elif match[1] != f"sun-valley-{mode}":
        problem = f"The app used the theme {match[1]}, not sun-valley-{mode}."
    else:
        return []
    return with_output(problem, output)


def check_smoke(exe: Path, timeout: float = TIMEOUT) -> list[str]:
    """Run ``exe --smoke`` with the light theme and with the dark theme. Return the problems."""
    for mode in ("light", "dark"):
        code, output = run_exe([str(exe), "--smoke"], timeout, env={THEME_ENV: mode})
        if problems := result_problems(code, output, timeout) or theme_problems(output, mode):
            return problems
        dnd = DND_LINE.search(output)
        print(
            f"The app passed the smoke test with the theme sun-valley-{mode}. "
            f"Drag and drop: {dnd[1] if dnd else 'unknown'}."
        )
    return []


def check(exe: Path, timeout: float = TIMEOUT) -> list[str]:
    """Run ``exe``, then run it on a test image. Return the problems, or an empty list."""
    if problems := check_smoke(exe, timeout):
        return problems
    with tempfile.TemporaryDirectory(prefix="pixheet-check-") as folder:
        src = Path(folder) / "test image.png"
        expected = make_image(src)
        start = time.monotonic()
        code, output = run_exe([str(exe), "--smoke", str(src)], timeout)
        seconds = time.monotonic() - start
        if problems := result_problems(code, output, timeout):
            return problems
        out = src.with_suffix(".xlsx")
        if not out.is_file():
            return [f"The app did not write {out.name}."]
        print(f"The app converted the test image in {seconds:.1f} s.")
        workbook = openpyxl.load_workbook(out)
        try:
            sheet = workbook.active
            size = (sheet.max_column, sheet.max_row)
            if size != (WIDTH, HEIGHT):
                return [f"The sheet has {size[0]} columns and {size[1]} rows."]
            problems = []
            for y in range(HEIGHT):
                for x in range(WIDTH):
                    rgb = sheet.cell(row=y + 1, column=x + 1).fill.fgColor.rgb
                    want = "FF" + bytes(expected.getpixel((x, y))).hex().upper()
                    if rgb != want:
                        problems.append(f"The cell ({x}, {y}) has {rgb}, not {want}.")
            if problems:
                return problems[:10]
        finally:
            workbook.close()
        return check_formats(exe, Path(folder) / FORMATS_FOLDER, timeout)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("exe", type=Path, help="the path of Pixheet.exe or of Pixheet")
    args = parser.parse_args()
    if not args.exe.is_file():
        print(f"check_exe: {args.exe} does not exist.", file=sys.stderr)
        return 1
    try:
        problems = check(args.exe.resolve())
    except CleanupError as e:
        problems = [str(e)]
    for problem in problems:
        print(f"check_exe: {problem}", file=sys.stderr)
    if not problems:
        print(f"check_exe: {args.exe} starts and converts all test images correctly.")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
