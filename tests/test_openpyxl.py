"""Round trip through openpyxl (a dev dependency only): fills match the source pixels."""

import random

import openpyxl
import pytest
from images import many_colors, noise
from PIL import Image
from sheets import hex_rgb

from pixheet import MAX_COLORS, convert, prepare


def small_noise() -> Image.Image:
    return noise(80, 60, 2_000, seed=4)


def transparent() -> Image.Image:
    return Image.new("RGBA", (30, 20), (10, 200, 30, 90))


@pytest.mark.parametrize(
    ("make", "compact"),
    [
        (small_noise, False),
        (small_noise, True),
        (transparent, False),
        (transparent, True),
        # Compact mode changes only the sheet, and most of the time here goes to
        # reading the 64k styles. The cases above cover compact sheets.
        (many_colors, False),
    ],
    ids=["noise", "noise-compact", "rgba", "rgba-compact", "many-colors"],
)
def test_fill_colors_match_pixels(tmp_path, make, compact):
    src = make()
    expected = prepare(src)
    w, h = expected.size
    out = tmp_path / "out.xlsx"
    result = convert(src, out, compact=compact)
    assert result.size == (w, h)
    assert result.colors <= MAX_COLORS

    wb = openpyxl.load_workbook(out)
    assert len(wb._cell_styles) == result.colors + 1  # plus the default style
    ws = wb.active
    assert (ws.max_column, ws.max_row) == (w, h)
    rng = random.Random(0)
    spots = [(0, 0), (w - 1, h - 1), (w - 1, 0), (0, h - 1)]
    spots += [(rng.randrange(w), rng.randrange(h)) for _ in range(200)]
    for x, y in spots:
        fill = ws.cell(row=y + 1, column=x + 1).fill
        assert fill.fill_type == "solid"
        assert fill.fgColor.rgb == f"FF{hex_rgb(expected.getpixel((x, y)))}"


def test_sheet_geometry(tmp_path, rgb_noise):
    out = tmp_path / "out.xlsx"
    convert(rgb_noise, out)
    ws = openpyxl.load_workbook(out).active
    assert ws.sheet_format.defaultRowHeight == 0.75
    assert ws.sheet_format.zeroHeight
    assert ws.row_dimensions[1].height == 0.75
    assert ws.column_dimensions["A"].width == pytest.approx(0.140625)
    assert ws.column_dimensions["A"].max == rgb_noise.size[0]
