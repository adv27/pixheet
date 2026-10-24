"""Tests of the desktop app. The widget tests skip when no display is available."""

import gc
import importlib.util
import os
import subprocess
import sys
import threading
import time
import types
import zipfile
from pathlib import Path

import pytest
from images import noise

from pixheet import _desktop, gui
from pixheet._plan import (
    IMAGE_SUFFIXES,
    Converted,
    Failed,
    Finished,
    Job,
    JobQueue,
    Progress,
    ReplaceQuestion,
    Started,
    Stopped,
    Tally,
    middle_ellipsis,
)
from pixheet.core import ConvertResult
from pixheet.core import convert as real_convert
from pixheet.gui import Options, Throttle, parse_options, run_queue

tk = pytest.importorskip("tkinter")
ttk = pytest.importorskip("tkinter.ttk")


@pytest.fixture(autouse=True)
def collect_tk_garbage():
    """Collect the Tk objects of a test on the main thread, at the end of the test.

    A Tk variable deletes itself in Tcl when Python collects it. The cyclic garbage
    collector can run on the worker thread of a later test. Then the call comes from the
    wrong thread. Tk waits for one second and then raises "main thread is not in main loop".
    """
    yield
    gc.collect()


@pytest.fixture(autouse=True)
def settings_folder(tmp_path, monkeypatch):
    """Point the settings file of the app at tmp_path, and remove ``PIXHEET_THEME``.

    Then no test reads or writes the settings of the user, and a theme that the user
    chose does not change a test.
    """
    folder = tmp_path / "config"
    monkeypatch.setenv("APPDATA", str(folder))
    monkeypatch.setenv("XDG_CONFIG_HOME", str(folder))
    monkeypatch.delenv(gui.THEME_ENV, raising=False)
    return folder


def image(path: Path, w: int = 4, h: int = 3) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    noise(w, h, min(w * h, 10)).save(path)
    return path


def destroyed(root) -> bool:
    try:
        return not root.winfo_exists()
    except tk.TclError:
        return True


def no_temp_files(folder: Path) -> bool:
    return not list(folder.glob(".pixheet-*.tmp"))


# --- Helpers without a display -------------------------------------------------------


def test_throttle_limits_updates_but_lets_the_first_and_last_through():
    now = [0.0]
    throttle = Throttle(0.05, clock=lambda: now[0])
    assert throttle(0, 10)  # the first
    now[0] = 0.01
    assert not throttle(1, 10)
    now[0] = 0.06
    assert throttle(2, 10)
    now[0] = 0.07
    assert not throttle(3, 10)
    assert throttle(10, 10)  # the last, at once after the one before


def test_the_worker_sends_about_four_progress_messages_each_second():
    # Tk draws the whole list again after each change of a cell, and that is slow.
    now = [0.0]
    throttle = Throttle(clock=lambda: now[0])
    assert throttle(0, 10)
    now[0] = 0.2
    assert not throttle(1, 10)
    now[0] = 0.25
    assert throttle(2, 10)


def test_image_suffixes_match_the_file_dialog_types():
    patterns = gui.IMAGE_TYPES[0][1].split()
    assert {pattern.removeprefix("*") for pattern in patterns} == IMAGE_SUFFIXES
    assert gui.image_types("win32") == gui.IMAGE_TYPES


def test_file_dialog_types_on_x11_ignore_the_case():
    [(name, patterns), (other, everything)] = gui.image_types("linux")
    assert (name, other, everything) == ("Images", "All files", "*")
    patterns = patterns.split()
    assert len(patterns) == len(IMAGE_SUFFIXES)
    assert "*.[pP][nN][gG]" in patterns
    assert "*.[tT][iI][fF][fF]" in patterns
    # Tcl "string match", which the Tk file dialog uses, accepts each case of a suffix.
    interp = tk.Tcl()
    for suffix in IMAGE_SUFFIXES:
        for name in (f"a{suffix}", f"a{suffix.upper()}", f"a{suffix.title()}"):
            assert any(interp.call("string", "match", p, name) for p in patterns), name
    assert not any(interp.call("string", "match", p, "a.txt") for p in patterns)


def test_parse_options():
    assert parse_options("1000", "white", False) == Options(1000, (255, 255, 255), False)
    assert parse_options(" 0 ", "none", True) == Options(0, None, True)
    assert parse_options("50", "#102030", False).background == (16, 32, 48)


@pytest.mark.parametrize(
    ("side", "background", "match"),
    [
        ("abc", "white", "maximum side"),
        ("-1", "white", "maximum side"),
        ("", "white", "maximum side"),
        ("100", "nope", "background"),
    ],
)
def test_parse_options_refuses_bad_text(side, background, match):
    with pytest.raises(ValueError, match=match):
        parse_options(side, background, False)


def run(jobs, cancel=None, interval=0.05):
    queue = JobQueue()
    queue.put(jobs)
    messages = []
    run_queue(queue, Options(), messages.append, cancel or threading.Event(), interval)
    return messages


def test_run_queue_messages(tmp_path):
    src = image(tmp_path / "in.png", 6, 5)
    dst = tmp_path / "out.xlsx"
    messages = run([Job(7, src, dst)], interval=0)

    assert messages[0] == Started(7)
    progress = [m for m in messages if isinstance(m, Progress)]
    assert [(m.job_id, m.done, m.total) for m in progress] == [(7, i, 5) for i in range(6)]
    converted = messages[-2]
    assert isinstance(converted, Converted)
    assert converted.job_id == 7
    assert converted.result.size == (6, 5)
    assert converted.size == dst.stat().st_size > 0
    assert messages[-1] == Finished(stopped=False)


def test_run_queue_continues_after_an_error(tmp_path):
    bad = tmp_path / "bad.png"
    bad.write_text("not an image")
    good = image(tmp_path / "good.png")
    messages = run([Job(1, bad, tmp_path / "bad.xlsx"), Job(2, good, tmp_path / "good.xlsx")])

    failed = [m for m in messages if isinstance(m, Failed)]
    assert failed == [Failed(1, gui.NOT_AN_IMAGE, "image")]  # no path with two backslashes
    assert messages[-1] == Finished(stopped=False)
    assert not (tmp_path / "bad.xlsx").exists()
    assert (tmp_path / "good.xlsx").exists()


def test_run_queue_gives_the_kind_of_a_write_error(tmp_path, monkeypatch):
    def convert(src, dst, **kwargs):
        raise PermissionError(13, "Access is denied", str(dst.with_name(".pixheet-0.tmp")))

    monkeypatch.setattr(gui, "convert", convert)
    messages = run([Job(1, image(tmp_path / "a.png"), tmp_path / "a.xlsx")])
    assert isinstance(messages[1], Failed)
    assert (messages[1].job_id, messages[1].kind) == (1, "write")


def test_run_queue_cancelled_before_the_start(tmp_path):
    cancel = threading.Event()
    cancel.set()
    messages = run([Job(1, image(tmp_path / "in.png"), tmp_path / "out.xlsx")], cancel)
    assert messages == [Finished(stopped=True)]
    assert not (tmp_path / "out.xlsx").exists()


def test_run_queue_cancelled_while_writing_keeps_the_old_file(tmp_path, monkeypatch):
    src = image(tmp_path / "in.png")
    dst = tmp_path / "out.xlsx"
    dst.write_bytes(b"old")
    cancel = threading.Event()

    def convert(*args, progress, **kwargs):
        def cancel_at_once(done, total):
            cancel.set()  # the user clicks Stop while convert writes the sheet
            progress(done, total)

        return real_convert(*args, progress=cancel_at_once, **kwargs)

    monkeypatch.setattr(gui, "convert", convert)
    messages = run([Job(1, src, dst), Job(2, src, tmp_path / "second.xlsx")], cancel)

    assert messages == [Started(1), Stopped(1), Finished(stopped=True)]
    assert dst.read_bytes() == b"old"
    assert not (tmp_path / "second.xlsx").exists()
    assert no_temp_files(tmp_path)


def test_run_queue_skips_one_job_and_continues(tmp_path, monkeypatch):
    first = Job(1, image(tmp_path / "a.png"), tmp_path / "a.xlsx")
    second = Job(2, image(tmp_path / "b.png"), tmp_path / "b.xlsx")
    removed = Job(3, image(tmp_path / "c.png"), tmp_path / "c.xlsx")
    removed.skip.set()  # the user removed the row after the worker took the job

    def convert(*args, progress, **kwargs):
        def skip_first(done, total):
            first.skip.set()  # the user clicks "Skip this file"
            progress(done, total)

        return real_convert(*args, progress=skip_first, **kwargs)

    monkeypatch.setattr(gui, "convert", convert)
    messages = run([first, removed, second])

    assert messages[:3] == [Started(1), Stopped(1), Stopped(3)]
    assert messages[3] == Started(2)
    assert isinstance(messages[-2], Converted)
    assert messages[-2].job_id == 2
    assert not first.dst.exists()
    assert not removed.dst.exists()
    assert messages[-1] == Finished(stopped=False)


def test_run_queue_takes_jobs_that_come_during_the_run(tmp_path):
    queue = JobQueue()
    later = Job(2, image(tmp_path / "b.png"), tmp_path / "b.xlsx")
    messages = []

    def post(message):
        messages.append(message)
        if message == Started(1):
            queue.put([later])  # a drop during the run

    queue.put([Job(1, image(tmp_path / "a.png"), tmp_path / "a.xlsx")])
    run_queue(queue, Options(), post, threading.Event())

    assert [m.job_id for m in messages if isinstance(m, Converted)] == [1, 2]
    assert messages[-1] == Finished(stopped=False)


def test_main_without_tkinter(monkeypatch, capsys):
    monkeypatch.setattr(gui, "tk", None)
    assert gui.main([]) == 1
    assert "tkinter is not available" in capsys.readouterr().err


@pytest.mark.parametrize("error", [ImportError, RuntimeError])
def test_enable_dnd_falls_back(monkeypatch, error):
    def fail():
        raise error("no tkdnd")

    monkeypatch.setattr(gui, "dnd_allowed", lambda platform=None: True)
    monkeypatch.setattr(gui, "_load_tkdnd", fail)
    assert gui.enable_dnd(object()) is None


def test_enable_dnd_does_not_load_tkdnd_when_it_can_stop_the_process(monkeypatch):
    monkeypatch.setattr(_desktop, "x11_supports_tkdnd", lambda load=None: False)
    monkeypatch.setattr(gui, "_load_tkdnd", lambda: pytest.fail("loads tkdnd"))
    assert gui.enable_dnd(object(), platform="linux") is None


def test_enable_dnd_can_be_turned_off(monkeypatch):
    monkeypatch.setenv(_desktop.DND_ENV, "0")
    monkeypatch.setattr(gui, "_load_tkdnd", lambda: pytest.fail("loads tkdnd"))
    assert gui.enable_dnd(object(), platform="win32") is None


@pytest.mark.parametrize(("platform", "opener"), [("linux", "xdg-open"), ("darwin", "open")])
def test_open_folder_does_not_wait_for_the_opener(tmp_path, monkeypatch, platform, opener):
    started = []
    real_popen = subprocess.Popen

    def popen(command, **kwargs):
        started.append((command, kwargs))
        # An opener that runs for some time, as xdg-open can.
        return real_popen([sys.executable, "-c", "import time; time.sleep(1)"], **kwargs)

    monkeypatch.setattr(_desktop.subprocess, "Popen", popen)
    t0 = time.monotonic()
    gui.open_folder(tmp_path / "a.xlsx", platform=platform)
    assert time.monotonic() - t0 < 0.5

    [(command, kwargs)] = started
    assert command == [opener, str(tmp_path)]
    assert kwargs["start_new_session"] is True
    assert kwargs["stdout"] == kwargs["stderr"] == kwargs["stdin"] == subprocess.DEVNULL
    waiter = next(t for t in threading.enumerate() if t.name == "pixheet-opener")
    waiter.join(10)
    assert not waiter.is_alive()


def test_open_folder_without_an_opener_raises(tmp_path, monkeypatch):
    def popen(command, **kwargs):
        raise FileNotFoundError(2, "No such file", command[0])

    monkeypatch.setattr(_desktop.subprocess, "Popen", popen)
    with pytest.raises(FileNotFoundError, match="xdg-open"):  # App shows it in a dialog
        gui.open_folder(tmp_path / "a.xlsx", platform="linux")


# --- The window ----------------------------------------------------------------------


def new_root():
    try:
        root = tk.Tk()
    except tk.TclError as e:
        pytest.skip(f"no display: {e}")
    root.withdraw()
    return root


class Harness:
    """An App on a hidden root, with dialogs that record their calls."""

    def __init__(self, answer="replace"):
        self.root = new_root()
        self.app = gui.App(self.root, dnd_type=None)
        self.errors = []
        self.questions = []
        self.folder = ""
        self.app.show_error = lambda title, message, **kw: self.errors.append(message)
        self.app.ask_replace = lambda question: self.questions.append(question) or answer
        self.app.ask_folder = lambda **kw: self.folder

    def drop(self, *paths):
        data = " ".join(f"{{{Path(p).as_posix()}}}" for p in paths)
        return self.app.on_drop(types.SimpleNamespace(data=data, action="copy"))

    def pump(self, until, timeout=20.0):
        deadline = time.monotonic() + timeout
        while not until():
            assert time.monotonic() < deadline, "timeout"
            self.root.update()
            time.sleep(0.005)

    def run(self, *paths):
        """Drop the paths, and wait until the run ends."""
        assert self.drop(*paths) == "copy"
        self.root.update()  # runs add_paths()
        self.wait()

    def wait(self):
        self.pump(lambda: not self.app.busy)

    def rows(self):
        """Return (name, status) for each row in the order of the list."""
        return [(row.src.name, gui.status_text(row)) for row in self.app.rows.values()]

    def states(self):
        return [row.state for row in self.app.rows.values()]

    def iid(self, name):
        return next(str(row.id) for row in self.app.rows.values() if row.src.name == name)

    def close(self):
        app = self.app
        app.cancel_event.set()
        for job in app.jobs_by_id.values():
            job.skip.set()
        if app.worker is not None:
            app.worker.join(10)
        try:
            app.dispose()
            self.root.destroy()
        except tk.TclError:  # the test destroyed the root
            pass


@pytest.fixture
def harness():
    made = []

    def make(**kwargs):
        made.append(Harness(**kwargs))
        return made[-1]

    yield make
    for h in made:
        h.close()


def gate_convert(monkeypatch):
    """Make convert wait at its first progress call until the test opens the gate."""
    gate = threading.Event()

    def convert(*args, progress, **kwargs):
        def wait_then_report(done, total):
            gate.wait(10)
            progress(done, total)

        return real_convert(*args, progress=wait_then_report, **kwargs)

    monkeypatch.setattr(gui, "convert", convert)
    return gate


def converting(h):
    return lambda: gui.CONVERTING in h.states()


# One file: the zone, as before the list.


def test_one_file_converts_in_the_zone(harness, tmp_path):
    h = harness()
    src = image(tmp_path / "my pics" / "one pic.png", 8, 6)

    h.run(src)

    dst = tmp_path / "my pics" / "one pic.xlsx"
    assert dst.exists()
    assert not h.app.list_mode
    assert h.app.zone.winfo_manager() == "grid"
    assert h.app.list_frame.winfo_manager() == ""
    assert h.app.title_var.get() == "Converted 1 file."
    status = h.app.status_var.get()
    assert status.startswith(str(dst))
    assert "8x6" in status
    assert "10 colors" in status
    assert h.app.progress_var.get() == 1.0
    assert h.app.open_button.instate(["!disabled"])
    assert h.app.open_button.cget("style") == "Accent.TButton"
    assert h.app.show_button.instate(["!disabled"])
    assert h.app.browse_button.instate(["!disabled"])
    assert h.app.stop_button.winfo_manager() == ""
    assert h.app.clear_button.winfo_manager() == ""
    assert h.root.title() == "Pixheet"
    assert h.errors == []
    assert no_temp_files(dst.parent)


def test_one_file_shows_its_progress_and_stop(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    assert h.drop(image(tmp_path / "IMG_2041.jpg")) == "copy"
    h.root.update()
    h.pump(lambda: h.app.title_var.get() == "File 1 of 1: IMG_2041.jpg")
    assert h.root.title() == "Pixheet - 1 of 1"
    assert h.app.status_var.get() == "Reading the image."
    assert h.app.stop_button.winfo_manager() == "grid"
    assert h.app.stop_button.cget("text") == "Stop"
    assert h.app.open_button.winfo_manager() == ""  # Stop only
    assert h.app.show_button.winfo_manager() == ""
    assert h.app.max_side_box.instate(["disabled"])
    assert h.app.save_box.instate(["disabled"])
    assert h.app.browse_button.instate(["!disabled"])  # a drop during the run is fine
    gate.set()
    h.wait()
    assert h.app.max_side_box.instate(["!disabled"])
    assert h.app.save_box.instate(["!disabled", "readonly"])


def test_one_file_that_fails_shows_the_error_in_red_without_a_dialog(harness, tmp_path):
    h = harness()
    bad = tmp_path / "IMG_2041.jpg"
    bad.write_text("not an image")

    h.run(bad)

    assert h.errors == []  # no message box
    assert h.app.title_var.get() == "Converted 0 of 1 file. 1 failed."
    assert h.app.status_var.get() == f"The app did not convert IMG_2041.jpg.\n{gui.NOT_AN_IMAGE}"
    assert "\\\\" not in h.app.status_var.get()
    assert str(h.app.status_label.cget("foreground")) == gui.ROW_COLORS["light"]["failed"]
    assert h.app.retry_button.cget("text") == "Retry failed (1)"
    assert h.app.open_button.winfo_manager() == ""  # no output to open


def test_drop_of_no_images_changes_nothing(harness, tmp_path):
    h = harness()
    (tmp_path / "notes.txt").write_text("x")
    assert h.drop(tmp_path / "notes.txt") == "copy"
    h.root.update()
    assert not h.app.busy
    assert h.app.rows == {}
    assert h.app.note_var.get() == (
        "The drop has no images. Ignored 1 file that is not an image: notes.txt."
    )
    assert h.app.note.winfo_manager() == "grid"


def test_bad_options_show_an_error_and_do_not_add_rows(harness, tmp_path):
    h = harness()
    h.app.max_side_var.set("abc")
    h.app.add_paths([image(tmp_path / "a.png")])
    assert not h.app.busy
    assert h.app.rows == {}
    assert len(h.errors) == 1
    assert "maximum side" in h.errors[0]


def test_options_go_to_convert(harness, tmp_path):
    h = harness()
    h.app.max_side_var.set("5")
    h.app.background_var.set("black")
    h.app.compact_var.set(True)
    src = image(tmp_path / "a.png", 20, 10)

    h.run(src)

    assert "5x2" in h.app.status_var.get()
    with zipfile.ZipFile(tmp_path / "a.xlsx") as z:
        assert 'r="A1"' not in z.read("xl/worksheets/sheet1.xml").decode()


# Two or more files: the list.


def five_from_three_folders(tmp_path):
    return [
        image(tmp_path / "Photos" / "2024" / "Da Lat" / "IMG_1042.jpg"),
        image(tmp_path / "Photos" / "2024" / "Da Lat" / "IMG_1043.jpg"),
        image(tmp_path / "Photos" / "2023" / "Beach" / "sunset.png"),
        image(tmp_path / "Photos" / "2023" / "Beach" / "pier.jpg"),
        image(tmp_path / "Scans" / "Family" / "grandma 1972.tif"),
    ]


def test_five_files_from_three_folders_show_the_list(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    sources = five_from_three_folders(tmp_path)
    h.drop(*sources)
    h.root.update()

    assert h.app.list_mode
    assert h.app.list_frame.winfo_manager() == "grid"
    assert h.app.zone.winfo_manager() == ""
    assert h.app.count_var.get() == "5 images from 3 folders"
    where = middle_ellipsis(str(tmp_path.parent), 40)
    assert h.app.tree.heading("folder", "text") == f"Folder (in {where})"
    folders = [h.app.tree.set(h.iid(src.name), "folder") for src in sources]
    base = tmp_path.name
    assert folders == [
        str(Path(base, "Photos", "2024", "Da Lat")),
        str(Path(base, "Photos", "2024", "Da Lat")),
        str(Path(base, "Photos", "2023", "Beach")),
        str(Path(base, "Photos", "2023", "Beach")),
        str(Path(base, "Scans", "Family")),
    ]
    assert tuple(h.app.tree.cget("displaycolumns")) == ("name", "folder", "status")
    h.pump(converting(h))
    assert h.app.title_var.get() == "File 1 of 5 · 0 done"
    assert h.root.title() == "Pixheet - 1 of 5"
    assert h.rows()[0] == ("IMG_1042.jpg", "Converting 0%")
    assert h.app.tree.item(h.iid("IMG_1042.jpg"), "tags") == ("converting",)
    assert h.rows()[1] == ("IMG_1043.jpg", "Waiting")
    assert h.app.clear_button.winfo_manager() == ""  # hidden during a run
    assert h.app.path_var.get() == h.app._path_line(sources[0], extra=" · Reading the image")

    gate.set()
    h.wait()
    assert h.app.title_var.get() == "Converted 5 files."
    assert h.root.title() == "Pixheet"
    assert all(status.startswith("✓ Done · ") for _name, status in h.rows())
    assert all(src.with_suffix(".xlsx").exists() for src in sources)
    assert h.app.clear_button.winfo_manager() == "grid"
    assert h.app.folders_button.cget("text") == "Show folders (3)"
    h.app._fill_folders_menu()
    labels = [h.app.folders_menu.entrycget(i, "label") for i in range(3)]
    assert labels[0] == f"{middle_ellipsis(str(sources[0].parent), 60)} (2)"
    assert labels[2] == f"{middle_ellipsis(str(sources[4].parent), 60)} (1)"
    assert h.app.open_button.instate(["disabled"])  # five outputs and no selection
    assert h.app.show_button.winfo_manager() == ""


def test_one_folder_hides_the_folder_column(harness, tmp_path):
    h = harness()
    folder = tmp_path / "Đà Lạt"
    h.run(image(folder / "a.png"), image(folder / "b.png"))
    assert h.app.count_var.get() == f"2 images in {middle_ellipsis(str(folder), 48)}"
    assert tuple(h.app.tree.cget("displaycolumns")) == ("name", "status")
    assert h.app.folder_open_button.winfo_manager() == "grid"
    assert h.app.folders_button.winfo_manager() == ""


def test_show_folders_opens_one_folder(harness, tmp_path, monkeypatch):
    opened = []
    monkeypatch.setattr(gui, "open_path", opened.append)
    h = harness()
    h.run(*five_from_three_folders(tmp_path))
    h.app._fill_folders_menu()
    h.app.folders_menu.invoke(1)
    assert opened == [tmp_path / "Photos" / "2023" / "Beach"]


def test_drop_during_a_run_adds_rows_at_the_end(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    first = [image(tmp_path / "a.png"), image(tmp_path / "b.png")]
    later = [image(tmp_path / "more" / "c.png"), image(tmp_path / "more" / "d.png")]
    h.drop(*first)
    h.root.update()
    h.pump(converting(h))

    assert h.drop(*later) == "copy"
    h.root.update()
    assert h.app.note_var.get() == "Added 2 images to the end of the list."
    assert [name for name, _status in h.rows()] == ["a.png", "b.png", "c.png", "d.png"]
    assert h.app.title_var.get().startswith("File 1 of 4")
    assert h.questions == []

    gate.set()
    h.wait()
    assert h.app.title_var.get() == "Converted 4 files."
    assert all(src.with_suffix(".xlsx").exists() for src in first + later)


def test_drop_during_a_run_uses_the_options_of_the_run(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    h.app.max_side_var.set("2")
    h.drop(image(tmp_path / "a.png", 8, 8), image(tmp_path / "b.png", 8, 8))
    h.root.update()
    h.pump(converting(h))
    h.app.max_side_var.set("abc")  # locked: the run does not read it
    h.drop(image(tmp_path / "c.png", 8, 8))
    h.root.update()
    gate.set()
    h.wait()
    assert h.errors == []
    assert sheet_size(tmp_path / "c.xlsx") == (2, 2)


def test_the_same_file_again(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    a = image(tmp_path / "a.png")
    b = image(tmp_path / "b.png")
    h.drop(a, b)
    h.root.update()
    h.pump(converting(h))

    h.drop(a)  # it converts now: ignore it
    h.root.update()
    assert h.app.note_var.get() == "1 image is already in the list."
    assert len(h.app.rows) == 2

    gate.set()
    h.wait()
    h.run(a)  # done: it goes back to the end of the queue, and the app asks first
    assert len(h.app.rows) == 2
    assert [name for name, _status in h.rows()] == ["b.png", "a.png"]
    assert [q.existing for q in h.questions] == [[tmp_path / "a.xlsx"]]
    assert h.questions[0].other_count == 0
    assert h.app.title_var.get() == "Converted 1 file."
    assert h.states() == ["done", "done"]


def test_the_same_file_again_uses_the_current_save_to(harness, tmp_path):
    h = harness()
    x = image(tmp_path / "x" / "photo.png")
    h.run(x)
    out = tmp_path / "out"
    out.mkdir()
    h.app.set_out_dir(out)

    y = image(tmp_path / "y" / "photo.png")
    h.run(x, y)  # the done row of x goes to the end of the queue, after the new y
    outputs = {row.src: row.dst for row in h.app.rows.values()}
    assert outputs == {x: out / "photo.xlsx", y: out / "photo (y).xlsx"}
    assert sorted(p.name for p in out.iterdir()) == ["photo (y).xlsx", "photo.xlsx"]
    assert h.questions == []


def test_a_drop_while_the_run_stops_waits(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    h.drop(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    h.root.update()
    h.pump(converting(h))
    h.app.stop_button.invoke()

    c = image(tmp_path / "c.png")
    h.drop(c)
    h.root.update()
    assert h.app.note_var.get() == "Added 1 image to the list. It waits, because the run stops."
    gate.set()
    h.wait()
    assert h.states() == ["stopped", "waiting", "waiting"]
    assert h.app.title_var.get() == "Stopped. Converted 0 of 2 files. 1 stopped."
    assert h.app.stop_button.cget("text") == "Convert 2"
    assert not c.with_suffix(".xlsx").exists()


def test_a_drop_while_the_run_stops_waits_when_the_worker_ends_in_the_question(
    harness, tmp_path, monkeypatch
):
    gate = gate_convert(monkeypatch)
    h = harness()
    h.drop(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    h.root.update()
    h.pump(converting(h))
    h.app.stop_button.invoke()
    c = image(tmp_path / "c.png")
    c.with_suffix(".xlsx").write_bytes(b"old")

    def ask(question):
        gate.set()
        h.pump(lambda: h.app.worker is None)  # the worker ends during the question
        return "replace"

    h.app.ask_replace = ask
    h.app.add_paths([c])
    h.root.update()
    assert h.app.note_var.get() == "Added 1 image to the list. It waits, because the run stops."
    assert h.app.worker is None
    assert h.states() == ["stopped", "waiting", "waiting"]
    assert h.app.title_var.get() == "Stopped. Converted 0 of 2 files. 1 stopped."
    assert h.app.stop_button.cget("text") == "Convert 2"
    assert c.with_suffix(".xlsx").read_bytes() == b"old"


def test_a_retry_while_the_run_stops_waits_when_the_worker_ends_in_the_question(
    harness, tmp_path, monkeypatch
):
    h = harness()
    bad = tmp_path / "bad.png"
    bad.write_bytes(b"not an image")
    h.run(bad)
    assert h.states() == ["failed"]
    gate = gate_convert(monkeypatch)
    h.drop(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    h.root.update()
    h.pump(converting(h))
    h.app.stop_button.invoke()
    bad.with_suffix(".xlsx").write_bytes(b"old")

    def ask(question):
        gate.set()
        h.pump(lambda: h.app.worker is None)  # the worker ends during the question
        return "replace"

    h.app.ask_replace = ask
    h.app.retry_failed()
    h.root.update()
    assert h.app.worker is None
    assert h.states() == ["waiting", "stopped", "waiting"]
    assert h.app.title_var.get() == "Stopped. Converted 0 of 2 files. 1 stopped."
    assert h.app.stop_button.cget("text") == "Convert 2"
    assert bad.with_suffix(".xlsx").read_bytes() == b"old"


def test_stop_then_continue(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    sources = [image(tmp_path / f"{name}.png") for name in "abc"]
    h.drop(*sources)
    h.root.update()
    h.pump(converting(h))

    h.app.stop_button.invoke()
    assert h.app.status_var.get() == "Stopping."
    gate.set()
    h.wait()

    assert h.states() == ["stopped", "waiting", "waiting"]
    assert h.rows()[0] == ("a.png", "\u2013 Stopped")
    assert h.app.title_var.get() == "Stopped. Converted 0 of 3 files. 1 stopped."
    assert h.app.stop_button.cget("text") == "Convert 2"
    assert h.app.stop_button.cget("style") == "Accent.TButton"
    assert h.app.retry_button.cget("text") == "Retry failed (1)"
    assert not any(src.with_suffix(".xlsx").exists() for src in sources)
    assert no_temp_files(tmp_path)

    h.app.stop_button.invoke()  # Convert 2
    h.wait()
    assert h.states() == ["stopped", "done", "done"]
    assert h.app.title_var.get() == "Converted 2 files."

    h.app.retry_button.invoke()
    h.wait()
    assert h.states() == ["done", "done", "done"]
    assert h.app.retry_button.winfo_manager() == ""


def test_skip_this_file_stops_one_file_and_the_run_continues(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    h.drop(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    h.root.update()
    h.pump(converting(h))
    h.app.tree.selection_set(h.iid("a.png"))
    menu = h.app.context_menu()
    assert menu.entrycget("Skip this file", "state") == "normal"
    menu.invoke("Skip this file")
    gate.set()
    h.wait()
    assert h.states() == ["stopped", "done"]
    assert h.app.title_var.get() == "Converted 1 of 2 files. 1 stopped."


def test_retry_failed_converts_the_failed_rows_again(harness, tmp_path):
    h = harness()
    bad = tmp_path / "bad.png"
    bad.write_text("not an image")
    good = image(tmp_path / "good.png")

    h.run(bad, good)
    assert h.states() == ["failed", "done"]
    assert h.rows()[0] == ("bad.png", "✕ Failed: not an image")
    assert h.app.tree.item(h.iid("bad.png"), "tags") == ("failed",)
    assert h.app.title_var.get() == "Converted 1 of 2 files. 1 failed."
    assert h.app.tree.selection() == (h.iid("bad.png"),)  # the first failed row
    assert h.app.path_var.get() == h.app._path_line(bad, bad.with_suffix(".xlsx"))
    assert h.app.status_var.get() == f"Failed: {gui.NOT_AN_IMAGE}"
    assert h.errors == []

    image(bad)  # the user fixes the file
    h.app.retry_button.invoke()
    h.wait()
    assert h.states() == ["done", "done"]
    assert h.app.title_var.get() == "Converted 1 file."


def test_failed_only_shows_only_the_failed_rows(harness, tmp_path):
    h = harness()
    bad = tmp_path / "b.png"
    bad.write_text("not an image")
    h.run(image(tmp_path / "a.png"), bad, image(tmp_path / "c.png"))
    button = h.app.failed_only_button
    assert button.winfo_manager() == "grid"
    assert button.cget("text") == "Failed only (1)"
    assert h.app.eta_label.winfo_manager() == ""

    button.invoke()
    assert h.app.tree.get_children() == (h.iid("b.png"),)
    button.invoke()
    assert h.app.tree.get_children() == tuple(h.iid(n) for n in ("a.png", "b.png", "c.png"))


def test_failed_only_ends_when_the_user_removes_the_failed_rows(harness, tmp_path):
    h = harness()
    bad = [tmp_path / "b.png", tmp_path / "d.png"]
    for path in bad:
        path.write_text("not an image")
    h.run(image(tmp_path / "a.png"), bad[0], image(tmp_path / "c.png"), bad[1])
    app = h.app
    button = app.failed_only_button
    button.invoke()
    assert button.cget("text") == "Failed only (2)"

    app.remove_rows([app.rows[int(h.iid("b.png"))]])
    assert app.failed_only
    assert button.cget("text") == "Failed only (1)"
    assert app.tree.get_children() == (h.iid("d.png"),)

    app.remove_rows([app.rows[int(h.iid("d.png"))]])
    assert not app.failed_only
    assert not app.failed_only_var.get()
    assert button.winfo_manager() == ""
    assert app.eta_label.winfo_manager() == "grid"
    assert app.tree.get_children() == (h.iid("a.png"), h.iid("c.png"))


def test_the_failed_only_count_follows_a_removal_when_the_filter_is_off(harness, tmp_path):
    h = harness()
    bad = [tmp_path / "b.png", tmp_path / "d.png"]
    for path in bad:
        path.write_text("not an image")
    h.run(image(tmp_path / "a.png"), bad[0], image(tmp_path / "c.png"), bad[1])
    app = h.app
    button = app.failed_only_button
    app.remove_rows([app.rows[int(h.iid("b.png"))]])
    assert button.cget("text") == "Failed only (1)"
    app.remove_rows([app.rows[int(h.iid("d.png"))]])
    assert button.winfo_manager() == ""
    assert app.eta_label.winfo_manager() == "grid"


def record_tree_writes(monkeypatch, tree):
    """Record each write of a cell or a tag of ``tree``. Reads are not in the list."""
    writes = []
    item, set_ = tree.item, tree.set

    def record_item(iid, option=None, **kw):
        if kw:
            writes.append(("item", iid, kw))
        return item(iid, option, **kw)

    def record_set(iid, column=None, value=None):
        if value is not None:
            writes.append(("set", iid, column, value))
        return set_(iid, column, value)

    monkeypatch.setattr(tree, "item", record_item)
    monkeypatch.setattr(tree, "set", record_set)
    return writes


def test_progress_writes_only_a_status_cell_that_changes(harness, monkeypatch, tmp_path):
    h = harness()
    gate = gate_convert(monkeypatch)
    assert h.drop(image(tmp_path / "a.png", 8, 6), image(tmp_path / "b.png", 8, 6)) == "copy"
    h.pump(converting(h))
    app = h.app
    row = next(row for row in app.rows.values() if row.state == gui.CONVERTING)
    iid = str(row.id)
    writes = record_tree_writes(monkeypatch, app.tree)

    app._handle(gui.Progress(row.id, 50, 100))
    assert writes == [("set", iid, "status", "Converting 50%")]
    writes.clear()
    app._handle(gui.Progress(row.id, 50, 100))  # the same text
    assert writes == []

    gate.set()
    h.wait()
    tags = [w for w in writes if w[0] == "item"]
    assert ("item", iid, {"tags": ("done",)}) in tags
    assert all(set(w[2]) == {"tags"} for w in tags)  # the name and the folder stay
    assert h.states() == [gui.DONE, gui.DONE]
    assert app.tree.item(iid, "values")[0] == row.src.name


def test_the_bar_moves_only_by_one_percent_or_more(harness, monkeypatch, tmp_path):
    h = harness()
    gate = gate_convert(monkeypatch)
    assert h.drop(image(tmp_path / "a.png", 8, 6), image(tmp_path / "b.png", 8, 6)) == "copy"
    h.pump(converting(h))
    app = h.app
    row = next(row for row in app.rows.values() if row.state == gui.CONVERTING)
    bar = []
    app.progress_var.trace_add("write", lambda *args: bar.append(app.progress_var.get()))

    app._handle(gui.Progress(row.id, 1, 100))  # 0.5 % of the run of 2 files
    assert bar == []
    app._handle(gui.Progress(row.id, 2, 100))
    assert bar == [0.01]
    app._handle(gui.Progress(row.id, 3, 100))
    assert bar == [0.01]

    gate.set()
    h.wait()
    assert app.progress_var.get() == 1.0


def test_the_poll_waits_longer_after_it_reads_messages(harness, monkeypatch, tmp_path):
    h = harness()
    gate = gate_convert(monkeypatch)
    assert h.drop(image(tmp_path / "a.png", 8, 6), image(tmp_path / "b.png", 8, 6)) == "copy"
    h.pump(converting(h))  # the worker now waits at the gate, and sends nothing
    app = h.app
    row = next(row for row in app.rows.values() if row.state == gui.CONVERTING)
    delays = []
    after = h.root.after

    def record_after(ms, func=None, *args):
        delays.append(ms)
        return after(ms, func, *args)

    monkeypatch.setattr(h.root, "after", record_after)
    for message in (None, gui.Progress(row.id, 50, 100)):
        h.root.after_cancel(app._poll_id)
        app._poll_id = None
        if message is not None:
            app.queue.put(message)
        app._poll()
    assert delays == [gui.POLL_MS, gui.BUSY_POLL_MS]
    assert gui.status_text(app.rows[row.id]) == "Converting 50%"

    gate.set()
    h.wait()
    assert h.states() == [gui.DONE, gui.DONE]


@pytest.mark.parametrize("answer", ["replace", "skip", "cancel"])
def test_outputs_that_exist_ask_once(harness, tmp_path, answer):
    h = harness(answer=answer)
    a = image(tmp_path / "x" / "a.png")
    b = image(tmp_path / "y" / "b.png")
    c = image(tmp_path / "y" / "c.png")
    for src in (a, b):
        src.with_suffix(".xlsx").write_bytes(b"old")

    h.run(a, b, c)

    [question] = h.questions
    assert question.existing == [a.with_suffix(".xlsx"), b.with_suffix(".xlsx")]
    assert question.other_count == 1
    assert question.heading == "2 output files exist"
    if answer == "replace":
        assert h.app.title_var.get() == "Converted 3 files."
        assert a.with_suffix(".xlsx").read_bytes() != b"old"
    elif answer == "skip":
        assert h.states() == ["skipped", "skipped", "done"]
        assert h.rows()[0] == ("a.png", "\u2013 Skipped: output exists")
        assert h.app.title_var.get() == "Converted 1 of 3 files. 2 skipped."
        assert a.with_suffix(".xlsx").read_bytes() == b"old"
        assert b.with_suffix(".xlsx").read_bytes() == b"old"
        assert c.with_suffix(".xlsx").exists()
    else:
        assert h.app.rows == {}
        assert h.app.note_var.get() == "The app did not add the dropped images."
        assert not c.with_suffix(".xlsx").exists()


def test_convert_again_replaces_a_skipped_output_without_a_question(harness, tmp_path):
    h = harness(answer="skip")
    a = image(tmp_path / "a.png")
    a.with_suffix(".xlsx").write_bytes(b"old")
    h.run(a, image(tmp_path / "b.png"))
    assert h.states() == ["skipped", "done"]

    h.app.tree.selection_set(h.iid("a.png"))
    menu = h.app.context_menu()
    assert menu.entrycget("Convert again", "state") == "normal"
    menu.invoke("Convert again")
    h.wait()
    assert len(h.questions) == 1
    assert h.states() == ["done", "done"]
    assert a.with_suffix(".xlsx").read_bytes() != b"old"


def test_two_types_in_one_folder_give_two_outputs(harness, tmp_path):
    h = harness()
    png = image(tmp_path / "photo.png")
    jpg = image(tmp_path / "photo.jpg")
    h.run(png, jpg)
    assert (tmp_path / "photo (png).xlsx").exists()
    assert (tmp_path / "photo (jpg).xlsx").exists()
    assert not (tmp_path / "photo.xlsx").exists()


def test_retry_keeps_the_tag_of_an_output(harness, tmp_path):
    h = harness()
    png = tmp_path / "photo.png"
    png.write_text("not an image")
    h.run(png, image(tmp_path / "photo.jpg"))
    assert h.states() == ["failed", "done"]

    image(png)  # the user fixes the file
    h.app.retry_button.invoke()
    h.wait()
    assert h.states() == ["done", "done"]
    assert sorted(p.name for p in tmp_path.glob("*.xlsx")) == [
        "photo (jpg).xlsx",
        "photo (png).xlsx",
    ]

    out = tmp_path / "out"
    out.mkdir()
    h.app.set_out_dir(out)
    h.run(png)  # another "Save to": a new output
    assert sorted(p.name for p in out.iterdir()) == ["photo.xlsx"]


def test_folder_drop_without_and_with_subfolders(harness, tmp_path):
    h = harness()
    beach = tmp_path / "Beach"
    for name in ("IMG_10.png", "IMG_2.png"):
        image(beach / name)
    image(beach / "day 1" / "IMG_5.png")
    image(beach / "day 2" / "IMG_6.png")
    (beach / "notes.txt").write_text("x")

    h.run(beach)
    assert [name for name, _status in h.rows()] == ["IMG_2.png", "IMG_10.png"]
    assert h.app.note_var.get() == (
        "Ignored 1 file that is not an image: notes.txt. "
        "The app did not open 2 subfolders of Beach."
    )
    assert h.app.note_button.winfo_manager() == "grid"
    assert h.app.note_button.cget("text") == "Include subfolders"
    assert h.app._note_id is None  # a note with a button stays

    h.app.note_button.invoke()
    h.root.update()
    h.wait()
    assert [name for name, _status in h.rows()] == [
        "IMG_2.png",
        "IMG_10.png",
        "IMG_5.png",
        "IMG_6.png",
    ]
    assert h.states() == ["done"] * 4
    assert h.questions == []  # the first two did not convert again
    assert h.app.note_var.get() == ""


def test_unwritable_output_folder_fails_in_its_row(harness, tmp_path, monkeypatch):
    def convert(src, dst, **kwargs):
        if src.name == "logo.png":
            temp = dst.with_name(".pixheet-0123456789abcdef.tmp")
            raise PermissionError(13, "Access is denied", str(temp))
        return real_convert(src, dst, **kwargs)

    monkeypatch.setattr(gui, "convert", convert)
    h = harness()
    logo = image(tmp_path / "Program Files" / "logo.png")
    h.run(logo, image(tmp_path / "ok.png"))

    assert h.rows()[0] == ("logo.png", "✕ Failed: cannot write here")
    assert h.app.tree.selection() == (h.iid("logo.png"),)
    system = "Windows" if sys.platform == "win32" else "The system"
    assert h.app.status_var.get() == (
        f"Failed: {system} does not let the app write in {middle_ellipsis(str(logo.parent), 40)}. "
        'Set "Save to" to a folder of your own, then click Retry failed.'
    )
    assert h.app.retry_button.cget("text") == "Retry failed (1)"

    # The user follows the advice. The text still names the folder that failed.
    status = h.app.status_var.get()
    h.app.set_out_dir(tmp_path / "Mine")
    assert h.app.tree.selection() == (h.iid("logo.png"),)
    assert h.app.status_var.get() == status
    copied = []
    h.app._copy = lambda lines: copied.extend(lines)  # not the clipboard of the user
    h.app.copy_errors()
    [line] = copied
    assert f"write in {logo.parent}." in line


def test_an_output_that_another_app_has_open_fails_in_its_row(harness, tmp_path, hold_open):
    h = harness()  # the answer to the replace question is "replace"
    a = image(tmp_path / "a.png", 8, 6)
    real_convert(a, tmp_path / "a.xlsx")
    close = hold_open(tmp_path / "a.xlsx")  # for example, Excel has the file open
    h.run(a, image(tmp_path / "b.png", 8, 6))

    assert h.rows()[0] == ("a.png", "✕ Failed: file is in use")
    assert h.states() == [gui.FAILED, gui.DONE]
    assert h.app.tree.selection() == (h.iid("a.png"),)
    assert h.app.status_var.get() == (
        "Failed: Windows cannot replace a.xlsx, because another app has the file open, "
        "for example Excel. Close the file in that app, then click Retry failed."
    )

    close()  # the user follows the advice
    h.app.retry_button.invoke()
    h.wait()
    assert h.states() == [gui.DONE, gui.DONE]


def test_a_long_output_path_fails_at_once(harness, tmp_path):
    h = harness()
    h.app.set_out_dir(tmp_path.joinpath(*["x" * 200] * 21))  # too long on all systems
    h.run(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    assert h.states() == ["failed", "failed"]
    assert h.rows()[0][1] == "✕ Failed: path too long"
    assert h.app.title_var.get() == "Converted 0 of 2 files. 2 failed."


def test_the_question_does_not_count_rows_that_fail_at_once(harness, tmp_path, monkeypatch):
    real_problem = gui.path_problem
    monkeypatch.setattr(
        gui, "path_problem", lambda dst: "too long" if dst.stem == "b" else real_problem(dst)
    )
    h = harness(answer="replace")
    a = image(tmp_path / "a.png")
    a.with_suffix(".xlsx").write_bytes(b"old")
    h.run(a, image(tmp_path / "b.png"))
    [question] = h.questions
    assert question.other_count == 0  # b fails at once: Skip converts nothing
    assert question.skip_label is None
    assert h.states() == ["done", "failed"]


def test_relative_paths_from_the_command_line(harness, tmp_path, monkeypatch):
    folder = tmp_path / "Photos"
    image(folder / "photo.png")
    image(folder / "sub" / "photo.png")
    monkeypatch.chdir(folder)
    h = harness()
    out = tmp_path / "out"
    out.mkdir()
    h.app.set_out_dir(out)

    h.app.add_paths([Path("photo.png"), Path("sub", "photo.png")])
    h.wait()
    assert [row.src for row in h.app.rows.values()] == [
        folder / "photo.png",
        folder / "sub" / "photo.png",
    ]
    assert sorted(p.name for p in out.iterdir()) == ["photo (Photos).xlsx", "photo (sub).xlsx"]


def test_save_to_one_folder_with_collisions(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    out = tmp_path / "out"
    out.mkdir()
    h.folder = str(out)
    h.app.save_var.set(gui.CHOOSE_FOLDER)
    h.app._on_save_to()
    assert h.app.out_dir == out
    assert h.app.save_var.get() == f"In {middle_ellipsis(str(out), 40)}"
    assert h.app.save_box.cget("values")[1] == h.app.save_var.get()

    b = image(tmp_path / "b" / "photo.png")
    c = image(tmp_path / "c" / "photo.png")
    h.drop(b, c, image(tmp_path / "c" / "other.png"))
    h.root.update()
    h.pump(converting(h))
    assert h.rows()[1] == ("photo.png", "Waiting · as photo (c).xlsx")
    h.app.tree.selection_set(h.iid("other.png"))
    h.root.update()
    other = tmp_path / "c" / "other.png"
    assert h.app.path_var.get() == h.app._path_line(other, out / "other.xlsx")

    gate.set()
    h.wait()
    assert sorted(p.name for p in out.iterdir()) == [
        "other.xlsx",
        "photo (b).xlsx",
        "photo (c).xlsx",
    ]


def test_save_to_cancel_keeps_the_old_value_and_next_to_each_image_goes_back(harness, tmp_path):
    h = harness()
    h.folder = str(tmp_path)
    h.app.save_var.set(gui.CHOOSE_FOLDER)
    h.app._on_save_to()
    h.folder = ""  # the user clicks Cancel
    h.app.save_var.set(gui.CHOOSE_FOLDER)
    h.app._on_save_to()
    assert h.app.out_dir == tmp_path
    assert h.app.save_var.get() == f"In {middle_ellipsis(str(tmp_path), 40)}"
    h.app.save_var.set(gui.NEXT_TO_IMAGE)
    h.app._on_save_to()
    assert h.app.out_dir is None
    assert list(h.app.save_box.cget("values")) == [gui.NEXT_TO_IMAGE, gui.CHOOSE_FOLDER]


def test_a_change_of_save_to_gives_new_outputs_to_waiting_rows(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    h.drop(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    h.root.update()
    h.pump(converting(h))
    h.app.cancel()
    gate.set()
    h.wait()
    out = tmp_path / "out"
    h.app.set_out_dir(out)
    assert [row.dst for row in h.app.rows.values()] == [out / "a.xlsx", out / "b.xlsx"]


def stop_after_the_first(h, tmp_path, monkeypatch):
    """Drop a.png and b.png, and stop the run in a.png. Then b.png waits."""
    gate = gate_convert(monkeypatch)
    h.drop(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    h.root.update()
    h.pump(converting(h))
    h.app.stop_button.invoke()
    gate.set()
    h.wait()
    assert h.states() == ["stopped", "waiting"]


def save_to_a_folder_with_b(h, tmp_path):
    """Set "Save to" to a folder that holds a b.xlsx of the user. Return that file."""
    out = tmp_path / "out"
    out.mkdir()
    (out / "b.xlsx").write_bytes(b"precious")
    h.app.set_out_dir(out)
    return out / "b.xlsx"


def test_convert_waiting_asks_about_outputs_that_exist(harness, tmp_path, monkeypatch):
    h = harness(answer="skip")
    stop_after_the_first(h, tmp_path, monkeypatch)
    precious = save_to_a_folder_with_b(h, tmp_path)

    assert h.app.stop_button.cget("text") == "Convert 1"
    h.app.stop_button.invoke()
    h.wait()
    assert [q.existing for q in h.questions] == [[precious]]
    assert precious.read_bytes() == b"precious"
    assert h.states() == ["stopped", "skipped"]


def test_convert_waiting_cancel_starts_no_run(harness, tmp_path, monkeypatch):
    h = harness(answer="cancel")
    stop_after_the_first(h, tmp_path, monkeypatch)
    precious = save_to_a_folder_with_b(h, tmp_path)

    h.app.convert_waiting()
    assert h.app.worker is None
    assert h.app.note_var.get() == gui.NOT_STARTED_TEXT
    assert h.states() == ["stopped", "waiting"]
    assert precious.read_bytes() == b"precious"


def test_a_drop_after_stop_asks_about_the_rows_that_wait(harness, tmp_path, monkeypatch):
    h = harness(answer="skip")
    stop_after_the_first(h, tmp_path, monkeypatch)
    precious = save_to_a_folder_with_b(h, tmp_path)

    h.run(image(tmp_path / "c.png"))
    [question] = h.questions
    assert question.existing == [precious]
    assert question.other_count == 1
    assert precious.read_bytes() == b"precious"
    assert h.states() == ["stopped", "skipped", "done"]
    assert (precious.parent / "c.xlsx").exists()


def test_retry_after_stop_asks_about_the_rows_that_wait(harness, tmp_path, monkeypatch):
    h = harness(answer="skip")
    stop_after_the_first(h, tmp_path, monkeypatch)
    precious = save_to_a_folder_with_b(h, tmp_path)

    h.app.retry_button.invoke()  # Retry failed (1): a.png, and the run takes b.png too
    h.wait()
    assert [q.existing for q in h.questions] == [[precious]]
    assert precious.read_bytes() == b"precious"
    assert h.states() == ["done", "skipped"]


def test_convert_waiting_checks_the_new_output_paths(harness, tmp_path, monkeypatch):
    h = harness()
    stop_after_the_first(h, tmp_path, monkeypatch)
    h.app.set_out_dir(tmp_path.joinpath(*["x" * 200] * 21))  # too long on all systems

    h.app.convert_waiting()
    h.wait()
    assert h.rows()[1] == ("b.png", "✕ Failed: path too long")


def test_clear_list_goes_back_to_the_zone(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    h.drop(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    h.root.update()
    h.app.clear_list()  # no effect during a run
    assert len(h.app.rows) == 2
    gate.set()
    h.wait()

    h.app.clear_button.invoke()
    assert h.app.rows == {}
    assert h.app.tree.get_children() == ()
    assert not h.app.list_mode
    assert h.app.zone.winfo_manager() == "grid"
    assert h.app.tally == Tally()
    assert h.app.title_var.get() == ""
    assert h.app.clear_button.winfo_manager() == ""

    h.run(image(tmp_path / "a.png"))  # alone again: the zone, and the plain name
    assert not h.app.list_mode
    assert (tmp_path / "a.xlsx").exists()


def test_keys_of_the_list(harness, tmp_path, monkeypatch):
    opened = []
    monkeypatch.setattr(gui, "open_path", opened.append)
    h = harness()
    sources = [image(tmp_path / f"{name}.png") for name in "abc"]
    h.run(*sources)
    tree = h.app.tree

    h.app._select_all()
    h.root.update()
    assert len(tree.selection()) == 3
    assert h.app.path_var.get() == "3 files selected."
    h.app.follow_again()  # Esc
    assert tree.selection() == ()
    assert h.app.follow

    tree.selection_set(h.iid("b.png"))
    h.app.activate()  # Enter
    assert opened == [tmp_path / "b.xlsx"]
    h.app.open_button.invoke()
    assert opened[-1] == tmp_path / "b.xlsx"

    h.app.remove_selected()  # Delete
    assert [name for name, _status in h.rows()] == ["a.png", "c.png"]
    assert h.app.list_mode  # the list stays until Clear list


def test_delete_during_a_run_keeps_the_converting_row(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    sources = [image(tmp_path / f"{name}.png") for name in "abc"]
    h.drop(*sources)
    h.root.update()
    h.pump(converting(h))
    h.app._select_all()
    menu = h.app.context_menu()
    assert menu.entrycget("Remove from list", "state") == "normal"
    h.app.remove_selected()
    assert [name for name, _status in h.rows()] == ["a.png"]
    gate.set()
    h.wait()
    assert (tmp_path / "a.xlsx").exists()
    assert not (tmp_path / "b.xlsx").exists()
    assert not (tmp_path / "c.xlsx").exists()
    assert h.app.title_var.get() == "Converted 1 file."


def test_context_menu_items_follow_the_selection(harness, tmp_path):
    h = harness()
    bad = tmp_path / "bad.png"
    bad.write_text("not an image")
    h.run(image(tmp_path / "a.png"), bad)

    def states(*names):
        h.app.tree.selection_set([h.iid(name) for name in names])
        menu = h.app.context_menu()
        labels = ["Open", "Show in folder", "Copy error", "Retry", "Convert again"]
        labels += ["Skip this file", "Remove from list", "Remove converted files"]
        return {label: menu.entrycget(label, "state") for label in labels}

    assert states("a.png") == {
        "Open": "normal",
        "Show in folder": "normal",
        "Copy error": "disabled",
        "Retry": "disabled",
        "Convert again": "disabled",
        "Skip this file": "disabled",
        "Remove from list": "normal",
        "Remove converted files": "normal",
    }
    failed = states("bad.png")
    assert (failed["Open"], failed["Copy error"], failed["Retry"]) == (
        "disabled",
        "normal",
        "normal",
    )
    # Record the calls. The system clipboard belongs to the user.
    clipboard = []
    h.root.clipboard_clear = clipboard.clear
    h.root.clipboard_append = clipboard.append
    h.app.copy_errors()
    assert clipboard == [f"{bad}: {gui.NOT_AN_IMAGE}"]
    h.app.copy_paths()
    assert clipboard == [str(bad)]


def test_the_context_menu_is_made_one_time(harness, tmp_path, monkeypatch):
    h = harness()
    h.run(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    posted = []
    monkeypatch.setattr(tk.Menu, "tk_popup", lambda menu, x, y, entry="": posted.append(menu))

    def menus():
        return [child for child in h.root.winfo_children() if isinstance(child, tk.Menu)]

    before = len(menus())
    commands = len(h.root.tk.call("info", "commands"))
    event = types.SimpleNamespace(x_root=10, y_root=10, y=5)
    for _ in range(20):
        h.app._on_right_click(event)
        h.app._on_menu_key()
    assert len(menus()) == max(before, 1)
    assert len(set(posted)) == 1
    assert len(h.root.tk.call("info", "commands")) <= commands + 12  # one menu at most
    assert posted[0].entrycget("Open", "state") in ("normal", "disabled")


def one_done_one_failed(h, tmp_path):
    """Convert a.png and bad.png. bad.png is not an image, so its row fails."""
    bad = tmp_path / "bad.png"
    bad.write_text("not an image")
    h.run(image(tmp_path / "a.png"), bad)
    assert h.states() == ["done", "failed"]
    return bad


def test_select_failed_files_and_remove_converted_files(harness, tmp_path):
    h = harness()
    one_done_one_failed(h, tmp_path)
    # The end of the run selects the failed row. Select the done row to see a change.
    h.app.tree.selection_set(h.iid("a.png"))
    h.app.select_failed()
    assert h.app.tree.selection() == (h.iid("bad.png"),)

    h.app.remove_done()
    assert [name for name, _status in h.rows()] == ["bad.png"]
    assert h.states() == ["failed"]
    assert (tmp_path / "a.xlsx").exists()  # the file stays, only the row goes

    h.app.clear_list()
    h.run(image(tmp_path / "c.png"))
    h.app.tree.selection_set(h.iid("c.png"))
    h.app.select_failed()  # no failed rows: the selection is empty
    assert h.app.tree.selection() == ()


def test_retry_in_the_menu_converts_only_the_failed_rows(harness, tmp_path):
    h = harness()
    bad = one_done_one_failed(h, tmp_path)
    image(bad)  # the user puts a real image in the place of the file
    (tmp_path / "a.xlsx").write_bytes(b"old")
    h.app._select_all()
    h.app.retry_selected()
    h.root.update()
    h.wait()
    assert h.states() == ["done", "done"]
    assert (tmp_path / "bad.xlsx").exists()
    assert (tmp_path / "a.xlsx").read_bytes() == b"old"  # a done row does not convert again
    assert h.questions == []


def test_open_error_is_shown(harness, tmp_path, monkeypatch):
    def fail(path):
        raise OSError("no association")

    monkeypatch.setattr(gui, "open_path", fail)
    h = harness()
    h.run(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    h.app.open_output()  # two done rows and no selection: nothing to open
    assert h.errors == []

    h.app.tree.selection_set(h.iid("a.png"))
    h.app.open_output()
    assert h.errors == ["Cannot open the file: no association"]
    assert h.app._dialogs == 0


def test_a_drop_while_the_window_closes_is_refused(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    h.drop(image(tmp_path / "a.png"))
    h.root.update()
    h.app.close()  # the window waits for the worker
    assert h.drop(image(tmp_path / "b.png")) == "refuse_drop"
    assert [name for name, _status in h.rows()] == ["a.png"]
    gate.set()
    h.pump(lambda: h.app.worker is None)
    assert destroyed(h.root)
    assert not (tmp_path / "b.xlsx").exists()


def test_note_button_without_an_action_closes_the_note(harness):
    h = harness()
    h.app.set_note("Converted 1 file.")
    h.app.run_note_action()
    assert h.app.note_var.get() == ""
    assert h.app.note.winfo_manager() == ""


def test_double_click_opens_a_done_row(harness, tmp_path, monkeypatch):
    opened = []
    monkeypatch.setattr(gui, "open_path", opened.append)
    h = harness()
    one_done_one_failed(h, tmp_path)
    # The list is not on the screen, so the test gives the row under the mouse.
    under_mouse = {10: h.iid("a.png"), 30: h.iid("bad.png"), 300: ""}
    monkeypatch.setattr(h.app.tree, "identify_row", under_mouse.get)

    h.app._on_double_click(types.SimpleNamespace(y=30))  # a failed row: no file to open
    assert h.app.tree.selection() == (h.iid("bad.png"),)
    assert opened == []
    h.app._on_double_click(types.SimpleNamespace(y=300))  # below the rows: no change
    assert h.app.tree.selection() == (h.iid("bad.png"),)
    h.app._on_double_click(types.SimpleNamespace(y=10))
    assert h.app.tree.selection() == (h.iid("a.png"),)
    assert opened == [tmp_path / "a.xlsx"]


def test_right_click_and_the_menu_key_open_the_context_menu(harness, tmp_path, monkeypatch):
    h = harness()
    one_done_one_failed(h, tmp_path)
    popups = []

    def popup(menu, x, y):
        popups.append(((x, y), menu.entrycget("Copy error", "state")))

    monkeypatch.setattr(h.app, "_popup", popup)
    under_mouse = {10: h.iid("a.png"), 30: h.iid("bad.png")}
    monkeypatch.setattr(h.app.tree, "identify_row", under_mouse.get)

    def click(y):
        return h.app._on_right_click(types.SimpleNamespace(y=y, x_root=500, y_root=600 + y))

    # A row outside the selection: the click selects it, and the menu fits that row.
    # The end of the run selects the failed row. Select the done row to see a change.
    h.app.tree.selection_set(h.iid("a.png"))
    assert click(30) == "break"
    assert h.app.tree.selection() == (h.iid("bad.png"),)
    # A row in the selection: the selection stays.
    h.app._select_all()
    assert click(10) == "break"
    assert len(h.app.tree.selection()) == 2
    assert popups == [((500, 630), "normal"), ((500, 610), "normal")]

    # The menu key opens the menu below the focus row, or near the corner of the list.
    popups.clear()
    tree = h.app.tree
    x, y = tree.winfo_rootx(), tree.winfo_rooty()
    tree.focus(h.iid("bad.png"))
    box = tree.bbox(h.iid("bad.png"))  # some systems give no box for a hidden list
    below = (x + box[0] + 20, y + box[1] + box[3]) if box else (x + 20, y + 20)
    assert h.app._on_menu_key() == "break"
    tree.focus("")
    assert h.app._on_menu_key() == "break"
    assert popups == [(below, "normal"), ((x + 20, y + 20), "normal")]


def test_popup_releases_the_grab_also_after_an_error():
    calls = []

    def tk_popup(x, y):
        calls.append(("popup", x, y))
        raise tk.TclError("cannot show the menu")

    menu = types.SimpleNamespace(tk_popup=tk_popup, grab_release=lambda: calls.append("release"))
    with pytest.raises(tk.TclError):
        gui.App._popup(None, menu, 5, 6)
    assert calls == [("popup", 5, 6), "release"]


def test_the_list_follows_the_converting_row(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    sources = [image(tmp_path / f"IMG_{i}.png", 2, 2) for i in range(30)]
    moves = []
    monkeypatch.setattr(h.app.tree, "yview_moveto", moves.append)
    h.drop(*sources)
    h.root.update()
    h.pump(converting(h))
    assert moves == [0.0]
    h.app._stop_follow()  # a click in the list
    gate.set()
    h.wait()
    assert moves == [0.0]


def test_a_tick_of_500_rows_does_not_count_the_rows(harness, tmp_path, monkeypatch):
    h = harness()
    app = h.app
    paths = [tmp_path / f"IMG_{i}.png" for i in range(500)]
    rows = [gui.Row(app._next_id(), path, path.with_suffix(".xlsx")) for path in paths]
    for row in rows:
        app._insert_row(row)
    app._refresh_list()
    app.worker = threading.current_thread()  # a run without a worker
    app.run_ids = {row.id for row in rows}
    app.run_tally = Tally(waiting=500)
    app._handle(Started(rows[250].id))
    monkeypatch.setattr(app, "_refresh_list", lambda: pytest.fail("counts the rows"))
    worst = 0.0
    for done in range(1, 101):
        t0 = time.perf_counter()
        app._handle(Progress(rows[250].id, done, 100))
        worst = max(worst, time.perf_counter() - t0)
    app.worker = None
    assert app.tree.set(str(rows[250].id), "status") == "Converting 100%"
    assert worst < 0.05, f"the worst tick took {worst * 1000:.1f} ms"


def test_open_and_show_in_folder_buttons(harness, tmp_path, monkeypatch):
    opened = []
    shown = []
    monkeypatch.setattr(gui, "open_path", opened.append)
    monkeypatch.setattr(gui, "open_folder", lambda path, select=False: shown.append((path, select)))
    h = harness()
    h.run(image(tmp_path / "a.png"))
    h.app.open_button.invoke()
    h.app.show_button.invoke()
    assert opened == [tmp_path / "a.xlsx"]
    assert shown == [(tmp_path / "a.xlsx", True)]


def test_open_folder_error_is_shown(harness, tmp_path, monkeypatch):
    def fail(path, select=False):
        raise OSError("no file manager")

    monkeypatch.setattr(gui, "open_folder", fail)
    h = harness()
    h.run(image(tmp_path / "a.png"))
    h.app.show_output()
    assert h.errors == ["Cannot open the folder: no file manager"]


@pytest.mark.parametrize("stop", ["cancel", "close"])
def test_stopping_text_stays_until_the_worker_stops(harness, tmp_path, monkeypatch, stop):
    gate = gate_convert(monkeypatch)
    h = harness()
    h.drop(image(tmp_path / "a.png"))
    h.root.update()  # runs add_paths(); the window reads the Started message later
    getattr(h.app, stop)()
    text = h.app.status_var.get()
    assert text.startswith("Stopping.")

    h.pump(lambda: h.app.title_var.get().startswith("File 1 of 1"))
    assert h.app.status_var.get() == text

    gate.set()
    h.pump(lambda: not h.app.busy)


def test_close_during_a_run_waits_for_the_worker(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    src = image(tmp_path / "a.png")
    dst = tmp_path / "a.xlsx"
    dst.write_bytes(b"old")
    h.drop(src)
    h.root.update()
    worker = h.app.worker

    h.app.close()
    assert h.app.closing
    assert h.root.winfo_exists()  # the window waits for the worker
    gate.set()

    h.pump(lambda: h.app.worker is None)
    assert not worker.is_alive()
    assert destroyed(h.root)
    assert dst.read_bytes() == b"old"
    assert no_temp_files(tmp_path)


def convert_threads() -> list[threading.Thread]:
    return [t for t in threading.enumerate() if t.name == "pixheet-convert"]


def drop_inside(h, path, seen, result):
    """Return a dialog that drops ``path`` and runs the event loop, as a modal dialog does."""

    def dialog(*args, **kwargs):
        seen.append(h.drop(path))
        h.root.update()  # runs an add_paths() that the drop scheduled
        seen.append(len(h.app.rows))
        return result

    return dialog


def test_drop_during_the_replace_dialog_is_refused(harness, tmp_path):
    h = harness()
    a = image(tmp_path / "a.png")
    a.with_suffix(".xlsx").write_bytes(b"old")
    b = image(tmp_path / "b.png")
    seen = []
    h.app.ask_replace = drop_inside(h, b, seen, "replace")

    h.run(a)

    assert seen == ["refuse_drop", 0]  # no row came in during the dialog
    assert h.app.note_var.get() == gui.DIALOG_DROP_TEXT
    assert h.app.title_var.get() == "Converted 1 file."
    assert a.with_suffix(".xlsx").read_bytes() != b"old"
    assert not (tmp_path / "b.xlsx").exists()
    assert convert_threads() == []


def test_drop_during_the_file_dialog_is_refused(harness, tmp_path):
    h = harness()
    a = image(tmp_path / "a.png")
    b = image(tmp_path / "b.png")
    seen = []
    h.app.ask_files = drop_inside(h, b, seen, (str(a),))

    h.app.browse()
    h.wait()

    assert seen == ["refuse_drop", 0]
    assert (tmp_path / "a.xlsx").exists()
    assert not (tmp_path / "b.xlsx").exists()


def test_add_folder_adds_the_images_of_the_folder(harness, tmp_path):
    h = harness()
    image(tmp_path / "pics" / "a.png")
    image(tmp_path / "pics" / "b.png")
    h.folder = str(tmp_path / "pics")
    h.app.folder_button.invoke()
    h.wait()
    assert h.states() == ["done", "done"]
    h.folder = ""
    h.app.browse_folder()  # Cancel
    assert len(h.app.rows) == 2


def test_drop_during_the_error_dialog_is_refused(harness, tmp_path):
    h = harness()
    h.app.max_side_var.set("abc")
    good = image(tmp_path / "good.png")
    seen = []
    h.app.show_error = drop_inside(h, good, seen, None)

    h.app.add_paths([image(tmp_path / "a.png")])

    assert seen == ["refuse_drop", 0]
    assert h.app.worker is None
    assert h.app._poll_id is None


def test_two_drops_before_idle_both_convert(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    a = image(tmp_path / "a.png")
    b = image(tmp_path / "b.png")
    assert h.drop(a) == "copy"
    assert h.drop(b) == "copy"
    h.root.update()
    assert len(convert_threads()) == 1

    gate.set()
    h.wait()
    assert (tmp_path / "a.xlsx").exists()
    assert (tmp_path / "b.xlsx").exists()


def test_close_during_a_dialog_closes_after_it(harness, tmp_path):
    h = harness()
    a = image(tmp_path / "a.png")
    a.with_suffix(".xlsx").write_bytes(b"old")
    alive = []

    def ask(question):
        h.app.close()
        alive.append(h.root.winfo_exists())  # the window waits for the dialog
        return "replace"

    h.app.ask_replace = ask
    h.app.add_paths([a])

    assert alive == [1]
    assert destroyed(h.root)
    assert h.app.worker is None
    assert a.with_suffix(".xlsx").read_bytes() == b"old"


def test_close_during_a_dialog_in_a_run_closes_when_both_end(harness, tmp_path, monkeypatch):
    gate = gate_convert(monkeypatch)
    h = harness()
    h.drop(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    h.root.update()
    h.pump(converting(h))
    c = image(tmp_path / "c.png")
    c.with_suffix(".xlsx").write_bytes(b"old")

    def ask(question):
        h.app.close()
        gate.set()
        h.pump(lambda: h.app.worker is None)  # the worker ends during the dialog
        assert h.root.winfo_exists()
        return "replace"

    h.app.ask_replace = ask
    h.app.add_paths([c])
    assert destroyed(h.root)
    assert c.with_suffix(".xlsx").read_bytes() == b"old"


def test_old_messages_do_not_reach_a_new_run(harness, tmp_path):
    h = harness()
    h.app.queue.put(Finished(stopped=True))  # a message that no run reads

    h.run(image(tmp_path / "a.png"))

    assert h.app.title_var.get() == "Converted 1 file."
    assert (tmp_path / "a.xlsx").exists()


def test_a_job_that_comes_when_the_worker_ends_starts_a_new_worker(harness, tmp_path):
    h = harness()
    h.run(image(tmp_path / "a.png"), image(tmp_path / "b.png"))
    app = h.app
    # The worker ends, and the window puts a job at the same time.
    app.worker = threading.Thread(target=lambda: None)
    app.worker.start()
    c = image(tmp_path / "c.png")
    row = gui.Row(app._next_id(), c, tmp_path / "c.xlsx")
    app._insert_row(row)
    app.run_ids.add(row.id)
    app.run_tally.add(None, row.state)
    app.jobs.put([Job(row.id, c, row.dst)])
    app._handle(Finished(stopped=False))
    assert app.worker is not None
    h.wait()
    assert (tmp_path / "c.xlsx").exists()


def test_stop_when_the_worker_ends_after_a_drop_stops_the_run(harness, tmp_path):
    h = harness()
    a = image(tmp_path / "a.png")
    b = image(tmp_path / "b.png")
    h.drop(a)
    h.root.update_idletasks()  # runs add_paths(), but not the poll
    h.app.worker.join(10)  # the worker took an empty queue and sent Finished(False)
    h.drop(b)
    h.root.update_idletasks()
    assert len(h.app.jobs) == 1
    h.app.stop_button.invoke()  # before the window reads Finished(False)

    h.wait()
    assert h.states() == ["done", "waiting"]
    assert not b.with_suffix(".xlsx").exists()
    assert h.app.title_var.get() == "Stopped. Converted 1 of 2 files."


def test_stop_after_the_last_file_is_no_stop(harness, tmp_path):
    h = harness()
    h.drop(image(tmp_path / "a.png"))
    h.root.update_idletasks()
    h.app.worker.join(10)
    h.app.stop_button.invoke()  # too late: nothing is left to stop
    h.wait()
    assert h.app.title_var.get() == "Converted 1 file."


def test_close_when_idle_destroys_the_root(harness):
    h = harness()
    h.app.close()
    assert destroyed(h.root)


def test_browse_converts_the_chosen_files(harness, tmp_path):
    h = harness()
    src = image(tmp_path / "a.png")
    h.app.ask_files = lambda **kw: (str(src),)
    h.app.browse()
    h.wait()
    assert (tmp_path / "a.xlsx").exists()


def test_browse_cancelled_does_nothing(harness):
    h = harness()
    h.app.ask_files = lambda **kw: ""
    h.app.browse()
    assert not h.app.busy


def test_window_without_dnd_says_so(harness):
    h = harness()
    assert not h.app.dnd
    assert h.app.status_var.get() == gui.NO_DND_TEXT
    assert h.app.zone_label.cget("text") == "Choose images to convert"
    assert h.app.or_label.winfo_manager() == ""  # no "or" above the buttons


def test_notes_without_a_button_close_after_some_time(harness, monkeypatch):
    monkeypatch.setattr(gui, "NOTE_MS", 10)
    h = harness()
    h.app.set_note("x")
    assert h.app.note.winfo_manager() == "grid"
    h.pump(lambda: h.app.note.winfo_manager() == "")
    assert h.app.note_var.get() == ""


def test_callback_error_is_shown_in_a_dialog(harness, capsys):
    h = harness()

    def boom():
        raise LookupError("boom")

    h.root.after_idle(boom)
    h.root.update()
    assert h.errors == ["An unexpected error occurred:\n\nLookupError('boom')"]
    assert "LookupError" in capsys.readouterr().err


def test_an_error_in_a_message_handler_does_not_stop_the_poll(harness, monkeypatch, tmp_path):
    h = harness()
    app = h.app
    gate = gate_convert(monkeypatch)
    handle = app._handle
    bad = []

    def handle_once_bad(message):
        if isinstance(message, gui.Progress) and not bad:
            bad.append(message)
            raise RuntimeError("boom in handler")
        handle(message)

    monkeypatch.setattr(app, "_handle", handle_once_bad)
    assert h.drop(image(tmp_path / "a.png", 8, 6), image(tmp_path / "b.png", 8, 6)) == "copy"
    h.pump(converting(h))
    gate.set()
    h.pump(lambda: h.errors)

    assert h.errors == ["An unexpected error occurred:\n\nRuntimeError('boom in handler')"]
    h.wait()  # the poll continues, and reads Finished
    assert h.states() == [gui.DONE, gui.DONE]
    app.close()
    assert destroyed(h.root)


def test_the_window_closes_after_an_error_in_a_message_handler(harness, monkeypatch, tmp_path):
    h = harness()
    app = h.app
    gate = gate_convert(monkeypatch)
    handle = app._handle

    def handle_once_bad(message):
        if isinstance(message, gui.Progress) and not h.errors:
            raise RuntimeError("boom in handler")
        handle(message)

    monkeypatch.setattr(app, "_handle", handle_once_bad)
    assert h.drop(image(tmp_path / "a.png", 8, 6), image(tmp_path / "b.png", 8, 6)) == "copy"
    h.pump(converting(h))
    gate.set()
    h.pump(lambda: h.errors)
    app.close()
    deadline = time.monotonic() + 20
    while not destroyed(h.root):
        assert time.monotonic() < deadline, "the window does not close"
        h.root.update()
        time.sleep(0.005)
    assert len(h.errors) == 1


def test_only_one_error_dialog_opens_at_a_time(harness, monkeypatch, tmp_path, capsys):
    h = harness()
    app = h.app
    handle = app._handle

    def handle_bad(message):
        if isinstance(message, gui.Progress):
            raise RuntimeError("boom in handler")
        handle(message)

    def show_error(title, message, **kwargs):
        # A dialog runs a nested event loop. The poll continues in it, and the next
        # messages also fail.
        h.errors.append(message)
        deadline = time.monotonic() + 20
        while app.worker is not None:
            assert time.monotonic() < deadline, "timeout"
            h.root.update()
            time.sleep(0.005)

    monkeypatch.setattr(app, "_handle", handle_bad)
    app.show_error = show_error
    h.run(image(tmp_path / "a.png", 8, 6), image(tmp_path / "b.png", 8, 6))

    assert h.errors == ["An unexpected error occurred:\n\nRuntimeError('boom in handler')"]
    assert capsys.readouterr().err.count("RuntimeError: boom in handler") >= 2
    assert h.states() == [gui.DONE, gui.DONE]


def test_main_smoke(monkeypatch, capsys):
    new_root().destroy()  # skip when there is no display
    assert gui.main(["--smoke"]) == 0
    assert "pixheet-gui: drag and drop" in capsys.readouterr().out


def test_main_adds_the_files_in_argv(monkeypatch):
    new_root().destroy()
    started = []
    monkeypatch.setattr(gui.App, "add_paths", lambda self, files: started.append(files))
    assert gui.main(["--smoke", "a b.png", "c.png"]) == 0
    assert started == [[Path("a b.png"), Path("c.png")]]


@pytest.mark.parametrize(
    ("value", "hidden"), [("1", True), (" 1 ", True), ("0", False), ("", False), ("yes", False)]
)
def test_hidden_window_reads_the_variable(value, hidden):
    assert gui.hidden_window({gui.HIDDEN_ENV: value}) is hidden
    assert gui.hidden_window({}) is False


class FakeWindow:
    """Records the calls of hide_window."""

    def __init__(self):
        self.calls = []

    def wm_geometry(self, value):
        self.calls.append(("geometry", value))

    def attributes(self, *args):
        self.calls.append(("attributes", *args))


@pytest.mark.parametrize(
    ("platform", "calls", "locked"),
    [
        ("win32", [("geometry", "+-32000+-32000"), ("attributes", "-toolwindow", True)], True),
        ("linux", [("geometry", "+-32000+-32000")], False),
    ],
)
def test_hide_window_places_the_window_off_the_screen(monkeypatch, platform, calls, locked):
    locks = []
    monkeypatch.setattr(gui, "_lock_foreground", lambda lock: locks.append(lock) or True)
    window = FakeWindow()
    assert gui.hide_window(window, platform=platform) is locked
    assert window.calls == calls
    assert locks == ([True] if locked else [])


def test_lock_foreground_does_nothing_outside_windows(monkeypatch):
    monkeypatch.setattr(gui.sys, "platform", "linux")
    assert gui._lock_foreground(True) is False


@pytest.mark.parametrize("hidden", [True, False])
def test_main_hides_the_window_in_the_hidden_mode(monkeypatch, hidden):
    new_root().destroy()
    monkeypatch.setenv(gui.HIDDEN_ENV, "1" if hidden else "0")
    locks, places = [], []
    monkeypatch.setattr(gui, "_lock_foreground", lambda lock: locks.append(lock) or True)
    hide = gui.hide_window

    def hide_window(root, platform=None):
        # conftest.py moves each test window to +-32000+-32000. Start from a different
        # place that is also off the screen, so that the test can see the move.
        root.wm_geometry("+-20000+-20000")
        result = hide(root, platform)
        root.after(100, lambda: places.append((root.winfo_rootx(), root.winfo_rooty())))
        return result

    monkeypatch.setattr(gui, "hide_window", hide_window)
    assert gui.main(["--smoke"]) == 0
    if hidden:
        [(x, y)] = places
        assert x < -30000
        assert y < -30000
        assert locks == ([True, False] if sys.platform == "win32" else [])
    else:
        assert places == locks == []


def sheet_size(path: Path) -> tuple[int, int]:
    """Return the number of columns and rows of the sheet in ``path``."""
    import openpyxl

    workbook = openpyxl.load_workbook(path)
    try:
        return workbook.active.max_column, workbook.active.max_row
    finally:
        workbook.close()


def test_main_smoke_waits_for_the_batch(tmp_path):
    new_root().destroy()
    src = image(tmp_path / "pic.png", 40, 30)
    assert gui.main(["--smoke", str(src)]) == 0
    assert sheet_size(src.with_suffix(".xlsx")) == (40, 30)
    assert no_temp_files(tmp_path)


def test_main_smoke_fails_when_a_file_fails(tmp_path, capsys):
    new_root().destroy()
    bad = tmp_path / "bad.png"
    bad.write_text("not an image")
    good = image(tmp_path / "good.png")
    assert gui.main(["--smoke", str(bad), str(good)]) == 1
    assert "bad.png" in capsys.readouterr().err
    assert good.with_suffix(".xlsx").is_file()


def test_main_smoke_does_not_replace_files(tmp_path):
    new_root().destroy()
    src = image(tmp_path / "pic.png")
    src.with_suffix(".xlsx").write_bytes(b"old")
    assert gui.main(["--smoke", str(src)]) == 0
    assert src.with_suffix(".xlsx").read_bytes() == b"old"


def test_main_smoke_of_a_built_app_needs_the_icon_and_dnd(tmp_path, monkeypatch, capsys):
    new_root().destroy()
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "_MEIPASS", str(tmp_path), raising=False)  # no icon here
    monkeypatch.setattr(gui, "enable_dnd", lambda root: None)
    monkeypatch.setattr(gui, "dnd_allowed", lambda: True)
    assert gui.main(["--smoke"]) == 1
    err = capsys.readouterr().err
    assert "The window icon is missing." in err
    assert "Drag and drop is not available." in err


def test_main_smoke_of_a_built_app_accepts_no_dnd_when_tkdnd_cannot_load(monkeypatch, capsys):
    new_root().destroy()
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(gui, "dnd_allowed", lambda *args: False)
    gui.main(["--smoke"])
    captured = capsys.readouterr()
    assert "pixheet-gui: drag and drop off" in captured.out
    assert "Drag and drop" not in captured.err


def test_missing_x11_dialogs():
    found = {"::tk::dialog::file::": "1", "::tk::dialog::file::chooseDir::": "1"}
    found["::tk::MessageBox"] = "0"

    def call(command, name):
        assert command == "auto_load"
        return found[name]

    assert gui.missing_x11_dialogs(call) == ["Tk cannot load the message box (msgbox.tcl)."]
    found["::tk::MessageBox"] = 1  # Tcl can give an int
    assert gui.missing_x11_dialogs(call) == []


def test_tk_finds_the_x11_dialogs():
    root = new_root()
    try:
        if root.tk.call("tk", "windowingsystem") != "x11":
            pytest.skip("Tk on this system uses the dialogs of the system")
        assert gui.missing_x11_dialogs(root.tk.call) == []
    finally:
        root.destroy()


def test_main_smoke_of_a_built_app_needs_the_theme(monkeypatch, capsys):
    new_root().destroy()
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(gui, "_load_sv_ttk", no_sv_ttk)
    assert gui.main(["--smoke"]) == 1
    assert "The Sun Valley theme is not available." in capsys.readouterr().err


# --- The theme -----------------------------------------------------------------------


def no_sv_ttk():
    raise ImportError("No module named 'sv_ttk'")


class FakeKey:
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def fake_winreg(value=None, error=None):
    """Return a fake winreg module. It has the value of AppsUseLightTheme, or raises."""

    def open_key(root, path):
        assert root == "HKCU"
        assert path == _desktop.PERSONALIZE_KEY
        if error is not None:
            raise error
        return FakeKey()

    def query(key, name):
        assert name == "AppsUseLightTheme"
        return value, 4

    return types.SimpleNamespace(HKEY_CURRENT_USER="HKCU", OpenKey=open_key, QueryValueEx=query)


@pytest.mark.parametrize(
    ("winreg", "theme"),
    [
        (fake_winreg(1), "light"),
        (fake_winreg(0), "dark"),
        (fake_winreg(error=FileNotFoundError(2, "The key does not exist")), "light"),
        (fake_winreg(error=PermissionError(5, "Access is denied")), "light"),
    ],
)
def test_system_theme_reads_the_windows_setting(monkeypatch, winreg, theme):
    monkeypatch.setitem(sys.modules, "winreg", winreg)
    assert gui.system_theme(platform="win32") == theme


def test_system_theme_without_winreg(monkeypatch):
    monkeypatch.setitem(sys.modules, "winreg", None)  # the import raises ImportError
    assert gui.system_theme(platform="win32") == "light"


def test_system_theme_on_other_systems_does_not_read_the_registry(monkeypatch):
    monkeypatch.setitem(sys.modules, "winreg", fake_winreg(0))
    monkeypatch.setattr(_desktop, "linux_theme", lambda: None)
    assert gui.system_theme(platform="linux") == "light"
    assert gui.system_theme(platform="darwin") == "light"


def test_system_theme_on_linux_asks_the_desktop(monkeypatch):
    monkeypatch.setattr(_desktop, "linux_theme", lambda: "dark")
    assert gui.system_theme(platform="linux") == "dark"


def test_follow_theme_reads_the_system_on_a_worker_thread_off_windows(monkeypatch):
    pytest.importorskip("sv_ttk")
    reads = []
    monkeypatch.setattr(gui, "system_theme", lambda: reads.append(1) or "dark")
    root = new_root()
    try:
        follower = gui.follow_theme(root, platform="linux")
        assert isinstance(follower.source, _desktop.LatestValue)
        assert follower.mode == "dark"
        assert reads == [1]
        follower.stop()
    finally:
        root.destroy()


@pytest.mark.parametrize(
    ("value", "theme"),
    [("light", "light"), (" Dark ", "dark"), ("System", "system"), ("", None), ("blue", None)],
)
def test_theme_override(value, theme):
    assert gui.theme_override({gui.THEME_ENV: value}) == theme
    assert gui.theme_override({}) is None


def test_pytest_does_not_capture_the_file_descriptors(pytestconfig):
    # With fd capture, Tcl keeps the closed handles of fd 1 and 2. Then "open" and "source"
    # fail sometimes for a file that a new handle with the same value points to.
    assert pytestconfig.getoption("capture") in ("sys", "no")


@pytest.mark.parametrize("mode", ["light", "dark"])
def test_apply_theme_uses_sun_valley(mode):
    pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        assert gui.apply_theme(root, mode) == f"sun-valley-{mode}"
        assert gui.theme_name(root) == f"sun-valley-{mode}"
    finally:
        root.destroy()


def test_disabled_accent_button_is_readable_in_the_light_theme():
    pytest.importorskip("sv_ttk")
    from tkinter import ttk

    root = new_root()
    try:
        gui.apply_theme(root, "dark")  # the change stays when the theme changes later
        gui.apply_theme(root, "light")
        colors = [(str(s), str(c)) for s, c in ttk.Style(root).map("Accent.TButton", "foreground")]
        assert ("disabled", gui.LIGHT_DISABLED_ACCENT_FG) in colors
    finally:
        root.destroy()


def test_apply_theme_scales_the_fonts_for_the_screen():
    pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        root.tk.call("tk", "scaling", 2.0)  # 144 dpi: 150 %
        gui.apply_theme(root, "light")
        assert root.tk.call("font", "configure", "SunValleyBodyFont", "-size") == -21
        gui.apply_theme(root, "dark")  # only once
        assert root.tk.call("font", "configure", "SunValleyBodyFont", "-size") == -21
    finally:
        root.destroy()


@pytest.mark.parametrize("error", [ImportError, RuntimeError, OSError, tk.TclError])
def test_apply_theme_keeps_the_ttk_theme_when_sv_ttk_fails(monkeypatch, capsys, error):
    def fail():
        raise error("no sv_ttk")

    monkeypatch.setattr(gui, "_load_sv_ttk", fail)
    root = new_root()
    try:
        before = gui.theme_name(root)
        assert gui.apply_theme(root, "dark") is None
        assert gui.theme_name(root) == before
    finally:
        root.destroy()
    # The cause is visible, for example in the output of a failed test.
    assert capsys.readouterr().err == "pixheet-gui: cannot use the Sun Valley theme: no sv_ttk\n"


def test_apply_theme_without_a_message_names_the_error_type(monkeypatch, capsys):
    def fail():
        raise ImportError

    monkeypatch.setattr(gui, "_load_sv_ttk", fail)
    root = new_root()
    try:
        assert gui.apply_theme(root, "dark") is None
    finally:
        root.destroy()
    assert "cannot use the Sun Valley theme: ImportError" in capsys.readouterr().err


def test_window_without_sv_ttk_has_the_old_look(monkeypatch, harness):
    monkeypatch.setattr(gui, "_load_sv_ttk", no_sv_ttk)
    h = harness()
    follower = gui.ThemeFollower(h.root, lambda: "dark")
    assert follower.theme is None
    assert follower._after_id is None
    assert str(h.app.zone.cget("relief")) == "groove"
    assert h.app.zone.cget("style") == ""


def test_main_without_sv_ttk_starts(monkeypatch, capsys):
    new_root().destroy()
    monkeypatch.setattr(gui, "_load_sv_ttk", no_sv_ttk)
    assert gui.main(["--smoke"]) == 0
    out = capsys.readouterr().out
    assert "pixheet-gui: theme " in out
    assert "sun-valley" not in out


@pytest.mark.parametrize("mode", ["light", "dark"])
def test_main_smoke_shows_the_theme(monkeypatch, capsys, mode):
    pytest.importorskip("sv_ttk")
    new_root().destroy()
    monkeypatch.setenv(gui.THEME_ENV, mode)
    assert gui.main(["--smoke"]) == 0
    assert f"pixheet-gui: theme sun-valley-{mode}\n" in capsys.readouterr().out


def test_window_in_sun_valley_uses_the_card_and_the_accent_button():
    pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        gui.apply_theme(root, "light")
        app = gui.App(root, dnd_type=None)
        assert app.zone.cget("style") == "Card.TFrame"
        assert app.browse_button.cget("style") == "Accent.TButton"
        assert str(app.zone_label.cget("font")) == "SunValleySubtitleFont"
        app.dispose()
    finally:
        root.destroy()


def test_disabled_spinbox_and_combobox_use_the_gray_text_of_the_style():
    pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        gui.apply_theme(root, "light")
        app = gui.App(root, dnd_type=None)
        boxes = (app.max_side_box, app.background_box)
        # tk_setPalette of sv-ttk gives a fixed text color, for example #1c1c1c. A fixed
        # color hides the gray text of the disabled state.
        for mode in ("dark", "light"):
            root.update_idletasks()
            assert [str(box.cget("foreground")) for box in boxes] == ["", ""], mode
            assert str(app.or_label.cget("foreground")) != ""  # tk_setPalette did run
            gui.apply_theme(root, mode)
        app.dispose()
    finally:
        root.destroy()


def test_row_colors_and_row_height_follow_the_theme():
    pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        gui.apply_theme(root, "light")
        app = gui.App(root, dnd_type=None)
        for mode in ("light", "dark", "light"):
            gui.apply_theme(root, mode)
            root.update()
            for tag in ("failed", "converting", "muted"):
                assert str(app.tree.tag_configure(tag, "foreground")) == gui.ROW_COLORS[mode][tag]
            assert str(app.note_label.cget("bg")) == gui.NOTE_COLORS[mode][0]
            assert (
                str(ttk.Style(root).lookup("Note.TFrame", "background"))
                == (gui.NOTE_COLORS[mode][0])
            )
            font = gui.tkfont.nametofont("SunValleyBodyFont", root=root)
            height = ttk.Style(root).lookup("Treeview", "rowheight")
            assert int(height) == font.metrics("linespace") + 8
        app.dispose()
    finally:
        root.destroy()


def test_combobox_list_uses_the_sun_valley_font():
    pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        gui.apply_theme(root, "dark")
        app = gui.App(root, dnd_type=None)
        popdown = root.tk.call("ttk::combobox::PopdownWindow", app.background_box)
        assert str(root.tk.call(f"{popdown}.f.l", "cget", "-font")) == "SunValleyBodyFont"
        app.dispose()
    finally:
        root.destroy()


class ThemeSource:
    """A fake system setting. It counts its reads."""

    def __init__(self, mode):
        self.mode = mode
        self.reads = 0

    def __call__(self):
        self.reads += 1
        return self.mode


def test_follower_changes_the_theme_when_the_system_changes(monkeypatch):
    pytest.importorskip("sv_ttk")
    bars = []
    monkeypatch.setattr(gui, "set_title_bar", lambda root, dark: bars.append(dark))
    applied = []
    real_apply = gui.apply_theme
    monkeypatch.setattr(
        gui, "apply_theme", lambda root, mode: applied.append(mode) or real_apply(root, mode)
    )
    root = new_root()
    try:
        source = ThemeSource("light")
        follower = gui.ThemeFollower(root, source, interval=10_000)
        assert follower.theme == "sun-valley-light"
        follower.poll()  # no change: no new theme and no new title bar
        assert applied == ["light"]
        assert bars == []
        source.mode = "dark"
        follower.poll()
        assert applied == ["light", "dark"]
        assert bars == [True]
        assert follower.mode == "dark"
        assert follower.theme == "sun-valley-dark"
        assert gui.theme_name(root) == "sun-valley-dark"
        assert source.reads == 3
    finally:
        root.destroy()


def test_follower_stops_when_the_theme_does_not_change(monkeypatch, capsys):
    pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        source = ThemeSource("light")
        follower = gui.ThemeFollower(root, source, interval=10_000)
        assert follower._after_id is not None
        monkeypatch.setattr(gui, "_load_sv_ttk", no_sv_ttk)
        source.mode = "dark"
        follower.poll()
        assert follower._after_id is None  # no new read, and no new error
        assert follower.mode == "light"
        assert follower.theme == "sun-valley-light"
        assert gui.theme_name(root) == "sun-valley-light"
        assert capsys.readouterr().err.count("cannot use the Sun Valley theme") == 1
    finally:
        root.destroy()


def test_follower_reads_the_source_on_a_timer():
    pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        source = ThemeSource("light")
        follower = gui.ThemeFollower(root, source, interval=1)
        deadline = time.monotonic() + 10
        while source.reads < 3:
            assert time.monotonic() < deadline, "timeout"
            root.update()
            time.sleep(0.005)
        follower.stop()
        reads = source.reads
        for _ in range(3):
            time.sleep(0.01)
            root.update()
        assert source.reads == reads
    finally:
        root.destroy()


def test_follower_sets_the_title_bar_when_the_window_shows(monkeypatch):
    pytest.importorskip("sv_ttk")
    bars = []
    monkeypatch.setattr(gui, "set_title_bar", lambda root, dark: bars.append(dark))
    root = new_root()
    try:
        gui.ThemeFollower(root, ThemeSource("dark"), choice="dark")
        frame = tk.Frame(root)
        frame.pack()
        root.update()
        assert bars == []  # the window is hidden, and a child does not count
        root.deiconify()
        root.update()
        assert bars == [True]
    finally:
        root.destroy()


@pytest.mark.skipif(sys.platform != "win32", reason="only the app on Windows hides its start")
def test_follower_hides_the_window_until_it_shows_in_the_theme(monkeypatch):
    pytest.importorskip("sv_ttk")
    bars = []
    monkeypatch.setattr(gui, "set_title_bar", lambda root, dark: bars.append(dark))
    root = new_root()
    try:
        gui.ThemeFollower(root, ThemeSource("dark"), choice="dark")
        assert float(root.attributes("-alpha")) == 0.0
        root.deiconify()
        root.update()
        assert bars == [True]  # the title bar is dark before the window becomes visible
        deadline = time.monotonic() + 10
        while float(root.attributes("-alpha")) < 1.0:
            assert time.monotonic() < deadline, "the window stays transparent"
            time.sleep(0.01)
            root.update()
        root.withdraw()
        root.deiconify()
        root.update()
        assert float(root.attributes("-alpha")) == 1.0  # only the first time
    finally:
        root.destroy()


def test_window_without_sv_ttk_is_not_hidden(monkeypatch):
    monkeypatch.setattr(gui, "_load_sv_ttk", no_sv_ttk)
    root = new_root()
    try:
        assert gui.ThemeFollower(root, ThemeSource("dark")).theme is None
        assert float(root.attributes("-alpha")) == 1.0  # no <Map> handler shows it later
    finally:
        root.destroy()


def test_follower_stops_when_the_window_closes():
    pytest.importorskip("sv_ttk")
    root = new_root()
    follower = gui.ThemeFollower(root, ThemeSource("dark"), interval=10_000)
    assert follower._after_id is not None
    root.destroy()
    assert follower._after_id is None


def test_follower_with_a_fixed_theme_does_not_read_the_system():
    pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        source = ThemeSource("light")
        follower = gui.ThemeFollower(root, source, choice="dark")
        assert follower.theme == "sun-valley-dark"
        assert follower._after_id is None
        assert source.reads == 0
    finally:
        root.destroy()


def test_follower_refuses_an_unknown_choice():
    with pytest.raises(ValueError, match="blue"):
        gui.ThemeFollower(None, ThemeSource("light"), choice="blue")


def test_follow_theme_uses_the_environment(monkeypatch):
    pytest.importorskip("sv_ttk")
    monkeypatch.setenv(gui.THEME_ENV, "light")
    monkeypatch.setattr(gui, "system_theme", lambda: pytest.fail("reads the system"))
    root = new_root()
    try:
        follower = gui.follow_theme(root)
        assert follower.theme == "sun-valley-light"
        assert follower._after_id is None
    finally:
        root.destroy()


# --- The theme toggle and the settings file ------------------------------------------


@pytest.mark.skipif(sys.platform != "win32", reason="a Windows path is not absolute on POSIX")
def test_settings_path_on_windows():
    path = gui.settings_path({"APPDATA": r"C:\Users\me\AppData\Roaming"}, platform="win32")
    assert path == Path(r"C:\Users\me\AppData\Roaming", "Pixheet", "settings.ini")
    assert gui.settings_path({}, platform="win32") is None
    assert gui.settings_path({"APPDATA": "relative"}, platform="win32") is None


@pytest.mark.skipif(sys.platform == "win32", reason="a POSIX path is not absolute on Windows")
def test_settings_path_on_other_systems():
    home = {"HOME": "/home/me"}
    assert gui.settings_path(home, "linux") == Path("/home/me/.config/pixheet/settings.ini")
    xdg = {"XDG_CONFIG_HOME": "/cfg", **home}
    assert gui.settings_path(xdg, "linux") == Path("/cfg/pixheet/settings.ini")
    relative = {"XDG_CONFIG_HOME": "cfg", **home}
    assert gui.settings_path(relative, "linux") == Path("/home/me/.config/pixheet/settings.ini")
    assert gui.settings_path({}, "darwin") is None


def test_settings_path_follows_the_environment(settings_folder):
    path = gui.settings_path()
    assert path is not None
    assert path.is_relative_to(settings_folder)


def test_theme_choice_without_a_file_is_system(tmp_path):
    assert gui.read_theme_choice(None) == "system"
    assert gui.read_theme_choice(tmp_path / "missing" / "settings.ini") == "system"
    assert gui.read_theme_choice(tmp_path) == "system"  # a folder, not a file


@pytest.mark.parametrize("choice", gui.THEME_CHOICES)
def test_theme_choice_round_trip(tmp_path, choice):
    path = tmp_path / "Pixheet" / "settings.ini"
    path.parent.mkdir()
    path.write_text("[other]\nkey = value\n", encoding="utf-8")
    assert gui.write_theme_choice(path, choice)
    assert gui.read_theme_choice(path) == choice
    assert "[other]\nkey = value\n" in path.read_text(encoding="utf-8")  # kept
    assert sorted(p.name for p in path.parent.iterdir()) == ["settings.ini"]  # no temp file


@pytest.mark.parametrize(
    "content",
    [
        b"theme = dark\n",  # no section
        b"[appearance]\ntheme = dark\ntheme = light\n",  # the same option two times
        b"[appearance]\ntheme = blue\n",
        b"[appearance]\ntheme = \xff\xfe\n",  # not UTF-8
        b"\x00\x01\x02",
    ],
)
def test_damaged_settings_give_system_and_a_new_choice_replaces_them(tmp_path, content):
    path = tmp_path / "settings.ini"
    path.write_bytes(content)
    assert gui.read_theme_choice(path) == "system"
    assert gui.write_theme_choice(path, "dark")
    assert gui.read_theme_choice(path) == "dark"


def test_write_theme_choice_does_not_raise(tmp_path, capsys):
    blocker = tmp_path / "file"
    blocker.write_text("", encoding="utf-8")
    path = blocker / "settings.ini"  # the folder is a file
    assert not gui.write_theme_choice(path, "dark")
    assert "pixheet-gui: cannot save the settings in " in capsys.readouterr().err
    assert gui.read_theme_choice(path) == "system"


def test_write_theme_choice_removes_its_temp_file(tmp_path, monkeypatch, capsys):
    path = tmp_path / "settings.ini"

    def fail(self, target):
        raise PermissionError("locked")

    monkeypatch.setattr(gui.Path, "replace", fail)
    assert not gui.write_theme_choice(path, "light")
    assert list(tmp_path.iterdir()) == []
    assert "locked" in capsys.readouterr().err


def test_follow_theme_uses_the_saved_choice(tmp_path, monkeypatch):
    pytest.importorskip("sv_ttk")
    monkeypatch.setattr(gui, "system_theme", lambda: pytest.fail("reads the system"))
    path = tmp_path / "settings.ini"
    gui.write_theme_choice(path, "dark")
    root = new_root()
    try:
        follower = gui.follow_theme(root, path)
        assert follower.choice == "dark"
        assert follower.theme == "sun-valley-dark"
        assert follower._after_id is None
    finally:
        root.destroy()


def test_follow_theme_without_a_saved_choice_follows_the_system(tmp_path, monkeypatch):
    pytest.importorskip("sv_ttk")
    monkeypatch.setattr(gui, "system_theme", lambda: "dark")
    root = new_root()
    try:
        follower = gui.follow_theme(root, tmp_path / "settings.ini")
        assert follower.choice == "system"
        assert follower.theme == "sun-valley-dark"
        assert follower._after_id is not None
    finally:
        root.destroy()


@pytest.mark.parametrize("value", ["system", "light"])
def test_environment_wins_over_the_saved_choice_and_the_toggle_shows_it(
    tmp_path, monkeypatch, value
):
    pytest.importorskip("sv_ttk")
    monkeypatch.setenv(gui.THEME_ENV, value)
    monkeypatch.setattr(gui, "system_theme", lambda: "light")
    path = tmp_path / "settings.ini"
    gui.write_theme_choice(path, "dark")
    root = new_root()
    try:
        follower = gui.follow_theme(root, path)
        app = gui.App(root, dnd_type=None, follower=follower, settings=path)
        assert follower.choice == value
        assert app.theme_var.get() == value
        assert gui.theme_name(root) == "sun-valley-light"
        assert gui.read_theme_choice(path) == "dark"  # the start does not write
        app.dispose()
    finally:
        root.destroy()


def test_choose_stops_and_starts_the_reads_of_the_system(monkeypatch):
    pytest.importorskip("sv_ttk")
    bars = []
    monkeypatch.setattr(gui, "set_title_bar", lambda root, dark: bars.append(dark))
    root = new_root()
    try:
        source = ThemeSource("light")
        follower = gui.ThemeFollower(root, source, interval=10_000)
        assert follower._after_id is not None
        assert follower.choose("dark")
        assert follower._after_id is None  # the reads stop
        assert (follower.choice, follower.mode, follower.theme) == (
            "dark",
            "dark",
            "sun-valley-dark",
        )
        assert gui.theme_name(root) == "sun-valley-dark"
        assert bars == [True]
        assert source.reads == 1
        assert follower.choose("light")
        assert follower._after_id is None
        assert bars == [True, False]
        source.mode = "dark"
        assert follower.choose("system")  # reads the system at once
        assert source.reads == 2
        assert follower._after_id is not None  # the reads start again
        assert (follower.choice, follower.mode) == ("system", "dark")
        assert gui.theme_name(root) == "sun-valley-dark"
        assert bars == [True, False, True]
        source.mode = "light"
        follower.poll()
        assert gui.theme_name(root) == "sun-valley-light"
        with pytest.raises(ValueError, match="blue"):
            follower.choose("blue")
    finally:
        root.destroy()


def test_choose_keeps_the_choice_when_the_theme_cannot_change(monkeypatch, capsys):
    pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        follower = gui.ThemeFollower(root, ThemeSource("light"), interval=10_000)
        monkeypatch.setattr(gui, "_load_sv_ttk", no_sv_ttk)
        assert not follower.choose("dark")
        assert (follower.choice, follower.mode) == ("system", "light")
        assert gui.theme_name(root) == "sun-valley-light"
    finally:
        root.destroy()
    without = gui.ThemeFollower(new_root(), ThemeSource("light"))
    try:
        assert without.theme is None
        assert not without.choose("dark")
    finally:
        without.root.destroy()
    capsys.readouterr()


def toggle_app(path, choice="system", source=None):
    """Return an App with the theme toggle, its root and its follower."""
    pytest.importorskip("sv_ttk")
    root = new_root()
    follower = gui.ThemeFollower(root, source or ThemeSource("light"), choice=choice)
    return gui.App(root, dnd_type=None, follower=follower, settings=path)


def test_toggle_is_at_the_right_of_the_header_row(tmp_path):
    app = toggle_app(tmp_path / "settings.ini")
    try:
        assert [b.cget("text") for b in app.theme_buttons] == ["System", "Light", "Dark"]
        assert {b.cget("style") for b in app.theme_buttons} == {"Toggle.TButton"}
        toggle = app.theme_buttons[0].master
        assert toggle.master is app.header
        assert toggle.grid_info()["column"] == 1
        assert toggle.grid_info()["sticky"] == "e"
        assert app.header.grid_info()["row"] == 0
        assert app.header.columnconfigure(0)["weight"] == 1  # the left part is free
        assert app.area.grid_info()["row"] == 2
        assert app.save_box.master.master is app.header
        assert app.theme_var.get() == "system"
    finally:
        app.dispose()
        app.root.destroy()


def test_toggle_changes_the_theme_and_saves_the_choice(tmp_path, monkeypatch):
    bars = []
    monkeypatch.setattr(gui, "set_title_bar", lambda root, dark: bars.append(dark))
    path = tmp_path / "Pixheet" / "settings.ini"
    source = ThemeSource("light")
    app = toggle_app(path, source=source)
    system, light, dark = app.theme_buttons
    try:
        dark.invoke()
        assert gui.theme_name(app.root) == "sun-valley-dark"
        assert app.follower._after_id is None
        assert bars == [True]
        assert gui.read_theme_choice(path) == "dark"
        light.invoke()
        assert gui.read_theme_choice(path) == "light"
        system.invoke()
        assert gui.read_theme_choice(path) == "system"
        assert app.follower._after_id is not None
        assert gui.theme_name(app.root) == "sun-valley-light"
    finally:
        app.dispose()
        app.root.destroy()


def test_toggle_does_not_write_without_a_choice(tmp_path):
    path = tmp_path / "settings.ini"
    app = toggle_app(path)
    try:
        app.root.update()
        app.theme_buttons[0].invoke()  # the choice that is in use
        app.choose_theme()
        assert not path.exists()
    finally:
        app.dispose()
        app.root.destroy()
    assert list(tmp_path.iterdir()) == []


def test_toggle_without_a_settings_file_does_not_save(tmp_path):
    app = toggle_app(None)
    try:
        app.theme_buttons[2].invoke()
        assert gui.theme_name(app.root) == "sun-valley-dark"
    finally:
        app.dispose()
        app.root.destroy()


def test_toggle_shows_the_old_choice_when_the_theme_cannot_change(tmp_path, monkeypatch):
    path = tmp_path / "settings.ini"
    app = toggle_app(path)
    try:
        monkeypatch.setattr(gui, "_load_sv_ttk", no_sv_ttk)
        app.theme_buttons[2].invoke()
        assert app.theme_var.get() == "system"
        assert not path.exists()
    finally:
        app.dispose()
        app.root.destroy()


def test_toggle_is_one_tab_stop_and_the_arrows_move_the_choice(tmp_path):
    app = toggle_app(tmp_path / "settings.ini", choice="light")
    system, light, dark = app.theme_buttons
    try:

        def stops():
            return [int(b.cget("takefocus")) for b in app.theme_buttons]

        assert stops() == [0, 1, 0]
        for button in app.theme_buttons:
            for key in ("<Left>", "<Right>", "<Up>", "<Down>"):
                assert button.bind(key), (button, key)
        assert app._step_theme(1) == "break"
        assert app.theme_var.get() == "dark"
        assert stops() == [0, 0, 1]
        app._step_theme(1)  # from the last choice to the first
        assert app.theme_var.get() == "system"
        app._step_theme(-1)
        assert app.theme_var.get() == "dark"
        assert gui.theme_name(app.root) == "sun-valley-dark"
        # The Tab key goes from the toggle to the next control, not to the next choice.
        assert str(dark.tk_focusNext()) != str(system)
        assert str(dark.tk_focusNext()) != str(light)
    finally:
        app.dispose()
        app.root.destroy()


def test_window_without_a_follower_or_sv_ttk_has_no_toggle(harness, monkeypatch):
    assert harness().app.theme_buttons == []
    monkeypatch.setattr(gui, "_load_sv_ttk", no_sv_ttk)
    root = new_root()
    try:
        follower = gui.ThemeFollower(root, ThemeSource("dark"))
        app = gui.App(root, dnd_type=None, follower=follower)
        assert app.follower is None
        assert app.theme_buttons == []
        app.dispose()
    finally:
        root.destroy()


def test_main_does_not_write_the_settings(settings_folder, monkeypatch, capsys):
    new_root().destroy()
    assert gui.main(["--smoke"]) == 0
    assert not settings_folder.exists()


def test_main_uses_the_saved_choice(settings_folder, capsys):
    pytest.importorskip("sv_ttk")
    new_root().destroy()
    gui.write_theme_choice(gui.settings_path(), "dark")
    assert gui.main(["--smoke"]) == 0
    assert "pixheet-gui: theme sun-valley-dark\n" in capsys.readouterr().out


def test_set_title_bar_does_nothing_on_other_systems():
    assert not gui.set_title_bar(object(), dark=True, platform="linux")


def test_set_title_bar_does_not_raise_without_a_window():
    class NoWindow:
        def winfo_id(self):
            raise tk.TclError("bad window path name")

    assert not gui.set_title_bar(NoWindow(), dark=True, platform="win32")


@pytest.mark.skipif(sys.platform != "win32", reason="the title bar color is a Windows feature")
def test_set_title_bar_on_a_real_window():
    root = new_root()
    try:
        assert not gui.set_title_bar(root, dark=True)  # Tk did not show the window yet
        root.deiconify()
        root.update()
        assert gui.set_title_bar(root, dark=True)
        assert gui.set_title_bar(root, dark=False)
    finally:
        root.destroy()


# --- The size of the window ----------------------------------------------------------


def test_minimum_size_is_at_least_the_requested_size(harness):
    root = harness().root
    width, height = root.minsize()
    assert width >= max(gui.MIN_WIDTH, root.winfo_reqwidth())
    assert height >= max(gui.MIN_HEIGHT, root.winfo_reqheight())


RESULT = ConvertResult(original_size=(867, 650), size=(867, 650), colors=9, cells=9, seconds=1.0)


def fill_the_list(app, tmp_path):
    """Put rows in three states and two output folders into ``app``, and a note."""
    a, b, c = tmp_path / "a" / "x.png", tmp_path / "b" / "y.png", tmp_path / "c" / "z.png"
    for row in (
        gui.Row(1, a, a.with_suffix(".xlsx"), state="done", result=RESULT, size=2_000_000),
        gui.Row(2, b, b.with_suffix(".xlsx"), state="failed", error="e", kind="other"),
        gui.Row(3, c, c.with_suffix(".xlsx"), state="done"),
    ):
        app._insert_row(row)
    app._refresh_all()
    app.set_note("Ignored 1 file that is not an image: notes.txt.", "Include subfolders", print)


@pytest.mark.parametrize("percent", [100, 125, 150])
@pytest.mark.parametrize("themed", [True, False])
@pytest.mark.parametrize("mode", ["zone", "list"])
def test_nothing_is_cut_off_at_the_minimum_size(percent, themed, mode, tmp_path):
    if themed:
        pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        root.tk.call("tk", "scaling", 96 * percent / 100 / 72)
        follower = None
        if themed:
            follower = gui.ThemeFollower(root, ThemeSource("light"), choice="light")
            assert follower.theme is not None
        app = gui.App(root, dnd_type=None, follower=follower)
        assert len(app.theme_buttons) == (3 if themed else 0)
        widgets = [
            app.save_box,
            app.max_side_box,
            app.background_box,
            app.compact_check,
            app.progressbar,
            *app.theme_buttons,
        ]
        if mode == "zone":
            app.worker = threading.current_thread()  # a run: the Stop button shows too
            app._set_running(True)
            widgets += [app.zone_label, app.browse_button, app.folder_button, app.stop_button]
        else:
            fill_the_list(app, tmp_path)
            assert app.list_mode
            app.tree.selection_set("1")
            app._on_select()
            widgets += [
                app.path_label,
                app.status_label,
                app.show_button,
                app.open_button,
                app.list_browse_button,
                app.list_folder_button,
                app.tree,
                app.note_button,
                app.note_close,
                app.retry_button,
                app.clear_button,
                app.folders_button,
            ]
        width, height = root.minsize()
        root.geometry(f"{width}x{height}")
        root.deiconify()
        root.update()
        for widget in widgets:
            assert widget.winfo_ismapped(), widget
            parent = widget.nametowidget(widget.winfo_parent())
            if widget is not app.tree:  # the list gets smaller
                assert widget.winfo_width() >= widget.winfo_reqwidth(), widget
                assert widget.winfo_height() >= widget.winfo_reqheight(), widget
            assert widget.winfo_x() + widget.winfo_width() <= parent.winfo_width(), widget
            assert widget.winfo_y() + widget.winfo_height() <= parent.winfo_height(), widget
        if mode == "list":  # the path line fits, and at least 3 rows of the list show
            font = app._label_font()
            assert font.measure(app.path_var.get()) <= app.path_label.winfo_width()
            rowheight = int(ttk.Style(root).lookup("Treeview", "rowheight") or 20)
            assert app.tree.winfo_height() >= 4 * rowheight, app.tree.winfo_height()
        app.worker = None
        app.dispose()
    finally:
        root.destroy()


@pytest.mark.parametrize("font_size", [None, 14])
@pytest.mark.parametrize("themed", [True, False])
def test_all_buttons_fit_at_the_minimum_size(themed, font_size, tmp_path):
    """The button row can show all its buttons. They must fit, also with a wide font.

    The Linux build in Debian 11 has wider fonts. There, "Open" went past the right edge.
    A larger default font gives the same problem on each system.
    """
    if themed:
        pytest.importorskip("sv_ttk")
    root = new_root()
    try:
        if font_size is not None:
            for name in ("TkDefaultFont", "TkTextFont", "TkHeadingFont", "TkMenuFont"):
                gui.tkfont.nametofont(name, root=root).configure(size=font_size)
        follower = None
        if themed:
            follower = gui.ThemeFollower(root, ThemeSource("light"), choice="light")
        app = gui.App(root, dnd_type=None, follower=follower)
        fill_the_list(app, tmp_path)
        d = tmp_path / "d" / "w.png"
        app._insert_row(gui.Row(4, d, d.with_suffix(".xlsx")))  # a waiting row
        app._refresh_all()
        app.tree.selection_set("1")
        app._on_select()
        buttons = [
            app.stop_button,
            app.retry_button,
            app.clear_button,
            app.folders_button,
            app.show_button,
            app.open_button,
        ]
        width, height = root.minsize()
        root.geometry(f"{width}x{height}")
        root.deiconify()
        root.update()
        bar = app.buttons
        for button in buttons:
            assert button.winfo_ismapped(), button
            assert button.winfo_width() >= button.winfo_reqwidth(), button
            assert button.winfo_x() + button.winfo_width() <= bar.winfo_width(), button
        assert bar.winfo_x() + bar.winfo_width() <= bar.master.winfo_width()
        app.dispose()
    finally:
        root.destroy()


# --- The window icon -----------------------------------------------------------------

REPO_ICON = Path(__file__).resolve().parents[1] / "packaging" / gui.ICON_NAME
REPO_PNG = REPO_ICON.with_name(gui.ICON_PNG)


def test_icon_path_from_a_checkout():
    assert gui.icon_path() == REPO_ICON
    assert REPO_ICON.is_file()


def test_icon_path_of_a_built_app(tmp_path, monkeypatch):
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "_MEIPASS", str(tmp_path), raising=False)
    assert gui.icon_path() is None
    (tmp_path / gui.ICON_NAME).write_bytes(REPO_ICON.read_bytes())
    assert gui.icon_path() == tmp_path / gui.ICON_NAME


def test_icon_path_of_an_installed_package(tmp_path, monkeypatch):
    monkeypatch.setattr(gui, "__file__", str(tmp_path / "site-packages" / "pixheet" / "gui.py"))
    assert gui.icon_path() is None


class IconRoot:
    """A root that records the calls of iconbitmap."""

    def __init__(self):
        self.calls = []

    def iconbitmap(self, **kwargs):
        self.calls.append(kwargs)


def test_set_icon_sets_the_default_icon_on_windows():
    root = IconRoot()
    assert gui.set_icon(root, platform="win32")
    assert root.calls == [{"default": str(REPO_ICON)}]


def test_set_icon_does_nothing_without_the_file(monkeypatch):
    root = IconRoot()
    monkeypatch.setattr(gui, "icon_path", lambda name=gui.ICON_NAME: None)
    assert not gui.set_icon(root, platform="win32")
    assert not gui.set_icon(root, platform="linux")
    assert root.calls == []


def test_icon_path_finds_the_png_icon():
    assert gui.icon_path(gui.ICON_PNG) == REPO_PNG
    assert REPO_PNG.is_file()


def test_set_icon_shows_the_png_icon_off_windows(tmp_path, monkeypatch):
    root = new_root()
    try:
        assert gui.set_icon(root, platform="linux")
        assert root._pixheet_icon.width() == root._pixheet_icon.height() == 128
        # The file goes away after the check: Tk raises an error, and the app continues.
        monkeypatch.setattr(gui, "icon_path", lambda name: tmp_path / name)
        assert not gui.set_icon(root, platform="linux")
    finally:
        root.destroy()


@pytest.mark.skipif(sys.platform != "win32", reason="Tk reads .ico files only on Windows")
def test_set_icon_on_a_real_window(tmp_path, monkeypatch):
    root = new_root()
    try:
        assert gui.set_icon(root)
        # The file goes away after the check: Tk raises an error, and the app continues.
        monkeypatch.setattr(gui, "icon_path", lambda: tmp_path / gui.ICON_NAME)
        assert not gui.set_icon(root)
    finally:
        root.destroy()


def test_window_has_the_icon(harness):
    assert harness().app.has_icon


def has_tkinterdnd2() -> bool:
    return importlib.util.find_spec("tkinterdnd2") is not None


SHOW_WINDOWS = os.environ.get("PIXHEET_TEST_SHOW_WINDOWS") == "1"


def run_python(*args: str) -> subprocess.CompletedProcess[str]:
    """Run Python in a new process. Tk loads tkdnd reliably only in a new interpreter.

    conftest.py does not move the windows of the new process off the screen. Thus the
    process gets the hidden mode of the app, except with PIXHEET_TEST_SHOW_WINDOWS=1.
    """
    env = os.environ | ({} if SHOW_WINDOWS else {gui.HIDDEN_ENV: "1"})
    return subprocess.run(
        [sys.executable, *args], capture_output=True, text=True, timeout=60, check=False, env=env
    )


def test_run_python_hides_the_window_of_the_app():
    result = run_python("-c", f"import os; print(os.environ.get({gui.HIDDEN_ENV!r}))")
    expected = os.environ.get(gui.HIDDEN_ENV) if SHOW_WINDOWS else "1"
    assert result.stdout.strip() == str(expected)


def test_python_m_pixheet_gui_smoke():
    new_root().destroy()
    result = run_python("-m", "pixheet.gui", "--smoke")
    assert result.returncode == 0, result.stderr
    state = "on" if has_tkinterdnd2() and _desktop.dnd_allowed() else "off"
    assert f"pixheet-gui: drag and drop {state}" in result.stdout


@pytest.mark.skipif(sys.platform != "win32", reason="the lock is a function of Windows")
def test_the_tests_hold_the_foreground_lock(foreground_lock):
    if not foreground_lock:
        pytest.skip("no foreground lock, for example with PIXHEET_TEST_SHOW_WINDOWS=1")
    # While this process holds the lock, a new process cannot take it. A click of the user
    # removes the lock, but renew_foreground_lock took it again just before this test.
    code = "import ctypes; print(ctypes.windll.user32.LockSetForegroundWindow(1))"
    result = run_python("-c", code)
    assert result.stdout.split() == ["0"], result.stderr


REGISTER_DROP = """
import tkinter as tk
from pixheet import gui

root = tk.Tk()
root.withdraw()
app = gui.App(root, dnd_type=gui.enable_dnd(root))
print(app.dnd)
print(repr(app.status_var.get()))
print(app.or_label.winfo_manager())
for widget in (app.body, app.zone):
    print("on_drop" in root.tk.call("bind", widget, "<<Drop>>"))
root.destroy()
"""


def test_real_tkdnd_registers_the_drop_zone():
    if not has_tkinterdnd2():
        pytest.skip("tkinterdnd2 is not installed")
    if not _desktop.dnd_allowed():
        pytest.skip("tkdnd cannot load here, for example with an old libX11 or libxcb")
    new_root().destroy()
    result = run_python("-c", REGISTER_DROP)
    assert result.returncode == 0, result.stderr
    assert result.stdout.split() == ["True", "''", "grid", "True", "True"]


# --- The replace dialog --------------------------------------------------------------


def existing_outputs(tmp_path, count):
    paths = []
    for i in range(count):
        path = tmp_path / ("Beach" if i % 2 else "Da Lat") / f"IMG_{i}.xlsx"
        path.parent.mkdir(exist_ok=True)
        path.write_bytes(b"old")
        paths.append(path)
    return paths


def descendants(widget):
    for child in widget.winfo_children():
        yield child
        yield from descendants(child)


@pytest.mark.parametrize("button", ["replace", "skip", "cancel", "close"])
def test_replace_dialog_answers(harness, tmp_path, button):
    h = harness()
    question = ReplaceQuestion(existing_outputs(tmp_path, 3), other_count=53)
    dialog = gui.ReplaceDialog(
        h.root, question, heading_font=h.app._big_font, strong_font=h.app._strong_font
    )
    texts = {key: b.cget("text") for key, b in dialog.buttons.items()}
    assert texts == {"replace": "Replace all", "skip": "Skip these 3", "cancel": "Cancel"}
    assert dialog.focus is dialog.buttons["skip"]  # Enter does not replace
    assert dialog.buttons["skip"].cget("style") == "Accent.TButton"
    rows = [dialog.tree.item(iid) for iid in dialog.tree.get_children()]
    assert [row["values"][0] for row in rows][:3] == [
        middle_ellipsis(str(tmp_path / "Da Lat"), 48),
        "      IMG_0.xlsx",
        "      IMG_2.xlsx",
    ]
    assert rows[0]["values"][1] == "2 files"
    assert rows[0]["tags"] == ["folder"]
    assert len(str(rows[1]["values"][1])) == len("2026-10-15 10:25")
    labels = [w.cget("text") for w in descendants(dialog.top) if isinstance(w, ttk.Label)]
    assert labels[:2] == [
        "3 output files exist",
        "They are in 2 folders. The app can replace them. "
        "Or it can skip these images and convert the other 53.",
    ]

    def click():
        if button == "close":
            h.root.tk.call(dialog.top.protocol("WM_DELETE_WINDOW"))
        else:
            dialog.buttons[button].invoke()

    h.root.after(50, click)
    assert dialog.run() == ("cancel" if button == "close" else button)
    assert destroyed(dialog.top)


def test_replace_dialog_for_one_file_without_other_images(harness, tmp_path):
    h = harness()
    [path] = existing_outputs(tmp_path, 1)
    question = ReplaceQuestion([path], other_count=0)
    dialog = gui.ReplaceDialog(
        h.root, question, heading_font=h.app._big_font, strong_font=h.app._strong_font
    )
    assert set(dialog.buttons) == {"replace", "cancel"}  # no Skip button
    assert dialog.buttons["replace"].cget("text") == "Replace"
    assert dialog.focus is dialog.buttons["cancel"]
    assert dialog.tree is None
    assert dialog.top.bind("<Escape>")
    h.root.after(50, dialog.buttons["replace"].invoke)
    assert dialog.run() == "replace"


def test_the_app_asks_with_the_replace_dialog(tmp_path):
    root = new_root()
    try:
        app = gui.App(root, dnd_type=None)
        a = image(tmp_path / "a.png")
        a.with_suffix(".xlsx").write_bytes(b"old")
        b = image(tmp_path / "b.png")
        seen = []

        def answer():
            top = next(w for w in root.winfo_children() if isinstance(w, tk.Toplevel))
            buttons = {w.cget("text"): w for w in descendants(top) if isinstance(w, ttk.Button)}
            seen.append(sorted(buttons))
            buttons["Skip this file"].invoke()

        root.after(100, answer)
        app.add_paths([a, b])
        deadline = time.monotonic() + 20
        while app.busy:
            assert time.monotonic() < deadline
            root.update()
        assert seen == [["Cancel", "Replace", "Skip this file"]]
        assert [row.state for row in app.rows.values()] == ["skipped", "done"]
        assert a.with_suffix(".xlsx").read_bytes() == b"old"
        app.dispose()
    finally:
        root.destroy()


def test_the_path_line_of_the_details_fits_the_window(harness, tmp_path):
    app = harness().app
    src = tmp_path / "Photos" / "2024" / "Da Lat" / "IMG_1073.jpg"
    app._line_width = 100_000
    assert app._path_line(src, src.with_suffix(".xlsx")) == f"{src}{gui.ARROW}IMG_1073.xlsx"
    out = tmp_path / "out" / "IMG_1073.xlsx"
    assert app._path_line(src, out) == f"{src}{gui.ARROW}{out}"
    assert app._path_line(src, extra=" · Writing row 3 of 9") == f"{src} · Writing row 3 of 9"
    font = app._label_font()
    for chars in (60, 90):
        app._line_width = font.measure("x" * chars)
        line = app._path_line(src, out)
        assert font.measure(line) <= app._line_width, line
        assert line.endswith("IMG_1073.xlsx")
        assert "…" in line
        line = app._path_line(src, extra=" · Writing row 3 of 9")
        assert font.measure(line) <= app._line_width, line
        assert line.endswith("IMG_1073.jpg · Writing row 3 of 9")
